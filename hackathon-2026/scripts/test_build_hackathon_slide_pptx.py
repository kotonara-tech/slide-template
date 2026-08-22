# -*- coding: utf-8 -*-
"""
test_build_hackathon_slide_pptx.py
==================================
ことテク2026 2days ハッカソン チーム発表スライド（9枚・表2枚・全スライドノート付き）を
生成する build_hackathon_slide_pptx.py の受け入れテスト。

想定する発表フォーマット: 発表8分 + デモ2分 + 質疑応答3分程度。

Red フェーズ: build_hackathon_slide_pptx.py と team-slide-template.md が存在しないため失敗する。
Green フェーズ: scripts/build_kotonara_pptx.py 系列と同一のテーマ・レイアウト生成ロジックのまま、
以下だけをこのテンプレートに合わせたスクリプトで通す:
  - 入力 md / 出力 pptx を argparse で受け取る（デフォルトは同梱テンプレート md）
  - {{プレースホルダー}} 残存チェック（--template-mode のときだけ免除）
  - N_SLIDES / LAYOUT_FOR_SLIDE / EXPECTED_TABLE_SLIDES の期待値

実行:
  python -m pytest -q hackathon-2026/scripts/test_build_hackathon_slide_pptx.py
"""

from __future__ import annotations

import importlib
import os
import re
import sys

import pytest

SCRIPTS_DIR = os.path.dirname(os.path.abspath(__file__))
HACKATHON_DIR = os.path.dirname(SCRIPTS_DIR)
TEMPLATE_MD_NAME = "team-slide-template.md"
SAMPLE_PPTX_NAME = "team-slide-template-kotonara.pptx"

if SCRIPTS_DIR not in sys.path:
    sys.path.insert(0, SCRIPTS_DIR)

# プレースホルダーは 1 行内に収まり、中に { } を含まない
PLACEHOLDER_RE = re.compile(r"\{\{[^{}\n]+\}\}")


@pytest.fixture(scope="module")
def mod():
    return importlib.import_module("build_hackathon_slide_pptx")


def all_deck_text(prs):
    """スライド本文＋表セル＋スピーカーノートの全テキスト。"""
    chunks = []
    for slide in prs.slides:
        for sh in slide.shapes:
            if sh.has_text_frame:
                chunks.append(sh.text_frame.text)
            if getattr(sh, "has_table", False) and sh.has_table:
                for row in sh.table.rows:
                    for cell in row.cells:
                        chunks.append(cell.text_frame.text)
        if slide.has_notes_slide:
            chunks.append(slide.notes_slide.notes_text_frame.text)
    return "\n".join(chunks)


def table_slide_numbers(prs):
    return [
        i for i, s in enumerate(prs.slides, start=1)
        if any(getattr(sh, "has_table", False) and sh.has_table for sh in s.shapes)
    ]


def test_module_constants(mod):
    """9枚構成・表2枚（5,7枚目）・表紙/クロージングのレイアウト割り当て。"""
    assert mod.N_SLIDES == 9
    assert mod.LAYOUT_FOR_SLIDE == {
        0: mod.L_COVER,
        4: mod.L_TABLE,
        6: mod.L_TABLE,
        8: mod.L_CLOSING,
    }
    assert mod.EXPECTED_TABLE_SLIDES == {5: "MVP", 7: "チーム"}
    assert os.path.basename(mod.TEMPLATE_MD) == TEMPLATE_MD_NAME
    assert os.path.isfile(mod.TEMPLATE_MD), "同梱テンプレート md が存在すること"
    assert os.path.dirname(mod.TEMPLATE_MD) == HACKATHON_DIR


def test_default_out_path(mod):
    """出力パス未指定時は <mdの場所>/<md名>-kotonara.pptx に導出される。"""
    p = os.path.join("x", "y", "my-team-slides.md")
    assert mod.default_out_path(p) == os.path.join("x", "y", "my-team-slides-kotonara.pptx")


def test_template_mode_build_passes(mod):
    """--template-mode で同梱テンプレート md から見本 pptx がビルドでき、検証も通る。"""
    assert mod.main(["--template-mode"]) == 0
    sample = os.path.join(HACKATHON_DIR, SAMPLE_PPTX_NAME)
    assert os.path.isfile(sample)

    from pptx import Presentation

    prs = Presentation(sample)
    assert len(prs.slides) == 9
    for i, slide in enumerate(prs.slides, start=1):
        assert slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip(), (
            f"スライド {i} にスピーカーノートが無い"
        )
    assert table_slide_numbers(prs) == [5, 7]

    assert prs.slides[0].shapes.title.text_frame.text == "{{プロダクト名}}"
    assert "ありがとうございました" in prs.slides[8].shapes.title.text_frame.text
    assert "ことテク2026" in (prs.core_properties.title or "")

    # 見本にはプレースホルダーがそのまま見えている（どこを書き換えるか一目で分かる）
    assert PLACEHOLDER_RE.search(all_deck_text(prs))

    # 時間配分（発表8分 + デモ2分 + 質疑3分）がノートに刻まれている
    assert "2分" in prs.slides[3].notes_slide.notes_text_frame.text  # 4枚目: デモ
    assert "3分" in prs.slides[8].notes_slide.notes_text_frame.text  # 9枚目: 質疑応答


def test_default_mode_rejects_placeholder_residue(mod, tmp_path):
    """{{...}} が残ったままの md は、通常モードではビルド失敗（0以外）になる。"""
    out = str(tmp_path / "reject.pptx")
    assert mod.main([mod.TEMPLATE_MD, "-o", out]) != 0


def test_filled_md_builds_clean(mod, tmp_path):
    """プレースホルダーを全部埋めた md は通常モードで ALL OK になり、残骸も無い。"""
    with open(mod.TEMPLATE_MD, encoding="utf-8") as fh:
        text = fh.read()
    filled = PLACEHOLDER_RE.sub("ダミー", text)
    assert not PLACEHOLDER_RE.search(filled)

    md = tmp_path / "team-x-slides.md"
    md.write_text(filled, encoding="utf-8")
    out = str(tmp_path / "team-x-kotonara.pptx")
    assert mod.main([str(md), "-o", out]) == 0

    from pptx import Presentation

    prs = Presentation(out)
    assert len(prs.slides) == 9
    assert table_slide_numbers(prs) == [5, 7]
    text_all = all_deck_text(prs)
    assert "{{" not in text_all and "}}" not in text_all
