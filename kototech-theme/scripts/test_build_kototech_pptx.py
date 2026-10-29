# -*- coding: utf-8 -*-
"""build_kototech_pptx.py の受け入れテスト。

実行（リポジトリルートで）:
  python -m pytest -q kototech-theme/scripts/test_build_kototech_pptx.py
"""
import os
import sys
import zipfile

from pptx import Presentation

SCRIPTS = os.path.dirname(os.path.abspath(__file__))
THEME_DIR = os.path.dirname(SCRIPTS)
sys.path.insert(0, SCRIPTS)

import build_kototech_pptx as K  # noqa: E402


def build(tmp_path, md, *extra):
    out = str(tmp_path / "out.pptx")
    code = K.main([os.path.join(THEME_DIR, md), "-o", out, *extra])
    return code, out


def test_template_builds_and_verifies(tmp_path):
    code, out = build(tmp_path, "template.md")
    assert code == 0
    prs = Presentation(out)
    assert [s.slide_layout.name for s in prs.slides] == [
        "表紙（タイトル）", "タイトルとコンテンツ", "表スライド", "セクション見出し", "クロージング（反転）"]


def test_theme_accent_and_empty_footer(tmp_path):
    _, out = build(tmp_path, "template.md")
    with zipfile.ZipFile(out) as z:
        theme = z.read("ppt/theme/theme1.xml").decode("utf-8")
    assert '<a:accent1><a:srgbClr val="4D9622"/></a:accent1>' in theme
    prs = Presentation(out)
    for layout in prs.slide_layouts:
        for sh in layout.placeholders:
            if sh.placeholder_format.type is not None and "FOOTER" in str(sh.placeholder_format.type):
                assert sh.text_frame.text == ""


def test_footer_option(tmp_path):
    _, out = build(tmp_path, "template.md", "--footer", "フッターの確認")
    prs = Presentation(out)
    texts = [sh.text_frame.text for sh in prs.slides[1].placeholders if sh.has_text_frame]
    assert "フッターの確認" in texts


def test_logos_live_in_layouts(tmp_path):
    _, out = build(tmp_path, "template.md")
    prs = Presentation(out)
    counts = {l.name: sum(1 for sh in l.shapes if sh.shape_type == 13) for l in prs.slide_layouts}
    assert counts == {"表紙（タイトル）": 1, "タイトルとコンテンツ": 1, "表スライド": 1,
                      "セクション見出し": 1, "クロージング（反転）": 0}
    for slide in prs.slides:
        assert not any(sh.shape_type == 13 for sh in slide.shapes)


def test_calendar_sample(tmp_path):
    code, out = build(tmp_path, "examples/calendar-sample.md")
    assert code == 0
    table = [sh for sh in Presentation(out).slides[1].shapes if sh.has_table][0].table
    assert len(table.columns) == 8
    assert table.cell(3, 6).text_frame.text == "19\n開催当日"


def test_placeholder_fails_without_draft(tmp_path):
    md = tmp_path / "p.md"
    md.write_text("# 表紙\n\n## {{サブタイトル}}\n\n---\n\n# おわり\n", encoding="utf-8")
    out = str(tmp_path / "p.pptx")
    assert K.main([str(md), "-o", out]) == 1
    assert K.main([str(md), "-o", out, "--draft"]) == 0
