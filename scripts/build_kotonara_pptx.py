# -*- coding: utf-8 -*-
"""
build_kotonara_pptx.py
======================
ことならテック (kotonara.tech) ブランドの PowerPoint デッキ／テンプレートを生成する。

生成物:
  1. kototech-aug-hackathon-kickoff-kotonara.pptx  … 本編19枚
  2. kotonara-brand-template.pptx                   … 同一テーマ／マスターの空テンプレート

配色・フォントを変えたいときは、下の CONFIG だけを書き換えて再実行すればよい。
theme1.xml のカラースキーム／フォントスキームとして書き出されるため、
PowerPoint の [デザイン] > [バリエーション] > [配色] / [フォント] からも差し替え可能。

実行:
  C:\\Users\\Ryuto\\AppData\\Local\\Programs\\Python\\Python312\\python.exe build_kotonara_pptx.py
"""

from __future__ import annotations

import copy
import io
import os
import re
import uuid
import zipfile
from xml.sax.saxutils import escape as xml_escape

from pptx import Presentation
from pptx.enum.text import MSO_ANCHOR
from pptx.oxml.ns import qn, nsdecls
from pptx.oxml import parse_xml
from pptx.util import Emu, Pt


# =============================================================================
# CONFIG  ―― ここだけ書き換えれば別配色・別フォントのデッキが生成される
# =============================================================================

CONFIG = {
    # テーマ名（PowerPoint の [デザイン]>[バリエーション] に表示される名前）
    "theme_name": "Kotonara",

    # --- テーマ12色（OOXML カラースキーム）。RRGGBB を # なしで指定 -------------
    "colors": {
        "dk1":      "222222",   # tx1  … 見出しの黒
        "lt1":      "FFFFFF",   # bg1  … 背景の白
        "dk2":      "333333",   # tx2  … 本文テキスト
        "lt2":      "EADECC",   # bg2  … ベージュ（カード背景）
        "accent1":  "1A832A",   # ブランドグリーン（メインアクセント）
        "accent2":  "4169E1",   # CTAブルー
        "accent3":  "337AB7",   # サブブルー
        "accent4":  "CCCC33",   # マスタード
        "accent5":  "EADECC",   # ベージュ
        "accent6":  "666666",   # 補助テキスト／罫線
        "hlink":    "4169E1",   # ハイパーリンク
        "folHlink": "337AB7",   # 訪問済リンク
    },

    # --- フォントスキーム（majorFont=見出し / minorFont=本文）------------------
    "fonts": {
        "major_latin": "Bahnschrift",   # 欧文見出し（Barlow 相当）
        "major_ea":    "游ゴシック",       # 和文見出し
        "minor_latin": "Segoe UI",      # 欧文本文（Lato 相当）
        "minor_ea":    "游ゴシック",       # 和文本文
    },

    # --- 文字サイズ（pt）-------------------------------------------------------
    "sizes": {
        "cover_title":     34,
        "cover_subtitle":  17,
        "cover_body":      16,
        "cover_catch":     15,
        "title":           30,
        "body":            18,
        "body_large":      20,
        "section_title":   32,
        "section_body":    18,
        "closing_title":   32,
        "closing_body":    18,
        "table_header":    15,
        "table_body":      14,
        "footer":          10,
    },

    # --- 定型文言 ---------------------------------------------------------------
    "footer_text": "ことテク2026 2daysハッカソン",
    "catch_copy":  "ひらめきを、カタチに。",

    # --- ドキュメントプロパティ -------------------------------------------------
    "doc_title":  "【ことテク2026】2days ハッカソン 事前オンラインキックオフ",
    "doc_author": "ことならテックプロジェクト",
}

# --- 入出力パス ---------------------------------------------------------------
BASE_DIR      = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
MD_PATH       = os.path.join(BASE_DIR, "kototech-aug-hackathon-kickoff-slides.md")
DECK_PATH     = os.path.join(BASE_DIR, "kototech-aug-hackathon-kickoff-kotonara.pptx")
TEMPLATE_PATH = os.path.join(BASE_DIR, "kotonara-brand-template.pptx")


# =============================================================================
# 名前空間・寸法定数
# =============================================================================

A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
NS_DECL = f'xmlns:a="{A_NS}" xmlns:r="{R_NS}" xmlns:p="{P_NS}"'

XML_HEAD = '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'

EMU_IN = 914400


def inch(v: float) -> int:
    return int(round(v * EMU_IN))


SW, SH = 12192000, 6858000            # 16:9 / 13.333in x 7.5in
ML = inch(0.833)                      # 左右マージン
CW = SW - 2 * ML                      # コンテンツ幅

TITLE_Y, TITLE_H = inch(0.525), inch(0.90)
BAR_Y, BAR_W, BAR_H = inch(1.541), inch(1.10), inch(0.05)
BODY_Y, BODY_H = inch(1.85), inch(4.50)
RULE_Y, RULE_H = inch(6.70), 9525
FTR_Y, FTR_H = inch(6.75), inch(0.30)
SN_W = inch(1.00)
SN_X = ML + CW - SN_W

TBL_Y = inch(2.00)
TBL_COL0 = inch(3.00)
TBL_COL1 = CW - TBL_COL0

FLD_SLIDENUM_GUID = "{" + "5D9C7C2E-6E44-4E4B-9C2E-7A1F0B3D5E11" + "}"
FLD_DATE_GUID = "{" + "1D3C7C2E-6E44-4E4B-9C2E-7A1F0B3D5E12" + "}"


def _sz(key: str) -> int:
    """CONFIG の pt サイズを OOXML の 1/100pt へ."""
    return int(CONFIG["sizes"][key] * 100)


# =============================================================================
# 低レベル XML ビルダ
# =============================================================================

SPTREE_HEAD = (
    "<p:spTree>"
    '<p:nvGrpSpPr><p:cNvPr id="1" name=""/><p:cNvGrpSpPr/><p:nvPr/></p:nvGrpSpPr>'
    '<p:grpSpPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="0" cy="0"/>'
    '<a:chOff x="0" y="0"/><a:chExt cx="0" cy="0"/></a:xfrm></p:grpSpPr>'
)
SPTREE_TAIL = "</p:spTree>"


def ph_sp(sid, name, ph, x, y, cx, cy, body_pr, lst_style="", paras="<a:p/>"):
    """プレースホルダー用 <p:sp>."""
    return (
        f'<p:sp><p:nvSpPr><p:cNvPr id="{sid}" name="{xml_escape(name)}"/>'
        f'<p:cNvSpPr><a:spLocks noGrp="1"/></p:cNvSpPr><p:nvPr>{ph}</p:nvPr></p:nvSpPr>'
        f'<p:spPr><a:xfrm><a:off x="{x}" y="{y}"/><a:ext cx="{cx}" cy="{cy}"/></a:xfrm>'
        f'<a:prstGeom prst="rect"><a:avLst/></a:prstGeom></p:spPr>'
        f"<p:txBody>{body_pr}{lst_style}{paras}</p:txBody></p:sp>"
    )


def rect_sp(sid, name, x, y, cx, cy, fill_xml, paras="<a:p/>", body_pr="<a:bodyPr/>"):
    """装飾用の矩形 <p:sp>（プレースホルダーではない）."""
    return (
        f'<p:sp><p:nvSpPr><p:cNvPr id="{sid}" name="{xml_escape(name)}"/>'
        f"<p:cNvSpPr/><p:nvPr/></p:nvSpPr>"
        f'<p:spPr><a:xfrm><a:off x="{x}" y="{y}"/><a:ext cx="{cx}" cy="{cy}"/></a:xfrm>'
        f'<a:prstGeom prst="rect"><a:avLst/></a:prstGeom>{fill_xml}<a:ln><a:noFill/></a:ln></p:spPr>'
        f"<p:txBody>{body_pr}<a:lstStyle/>{paras}</p:txBody></p:sp>"
    )


def solid(scheme, alpha=None):
    inner = f'<a:alpha val="{alpha}"/>' if alpha is not None else ""
    if inner:
        return f'<a:solidFill><a:schemeClr val="{scheme}">{inner}</a:schemeClr></a:solidFill>'
    return f'<a:solidFill><a:schemeClr val="{scheme}"/></a:solidFill>'


def def_rpr(sz, scheme, b=0, major=False, alpha=None, spc=None):
    pre = "mj" if major else "mn"
    spc_attr = f' spc="{spc}"' if spc is not None else ""
    return (
        f'<a:defRPr sz="{sz}" b="{b}" kern="1200"{spc_attr}>{solid(scheme, alpha)}'
        f'<a:latin typeface="+{pre}-lt"/><a:ea typeface="+{pre}-ea"/>'
        f'<a:cs typeface="+{pre}-cs"/></a:defRPr>'
    )


def lst_style_one(sz, scheme, b=0, major=False, algn="l", alpha=None,
                  bullet=True, marL=342900, spc_bef=900, spc=None, ln_spc=112000):
    """単一レベルの <a:lstStyle>（レイアウト側の書式上書き用）."""
    if bullet:
        bu = (
            f'<a:buClr><a:schemeClr val="accent1"/></a:buClr><a:buSzPct val="78000"/>'
            f'<a:buFont typeface="Arial" panose="020B0604020202020204" pitchFamily="34" charset="0"/>'
            f'<a:buChar char="\u25aa"/>'
        )
        indent = -marL
    else:
        bu = "<a:buNone/>"
        marL, indent = 0, 0
    return (
        f'<a:lstStyle><a:lvl1pPr marL="{marL}" indent="{indent}" algn="{algn}">'
        f'<a:lnSpc><a:spcPct val="{ln_spc}"/></a:lnSpc>'
        f'<a:spcBef><a:spcPts val="{spc_bef}"/></a:spcBef>{bu}'
        f"{def_rpr(sz, scheme, b, major, alpha, spc)}</a:lvl1pPr>"
        f'<a:lvl2pPr marL="{marL + 400050}" indent="{-285750}" algn="{algn}">'
        f'<a:lnSpc><a:spcPct val="{ln_spc}"/></a:lnSpc>'
        f'<a:spcBef><a:spcPts val="600"/></a:spcBef>'
        f'<a:buClr><a:schemeClr val="accent6"/></a:buClr><a:buSzPct val="90000"/>'
        f'<a:buFont typeface="Arial" panose="020B0604020202020204" pitchFamily="34" charset="0"/>'
        f'<a:buChar char="\u2013"/>'
        f"{def_rpr(max(sz - 200, 1000), 'accent6', 0, major)}</a:lvl2pPr></a:lstStyle>"
    )


BODY_PR_TIGHT = '<a:bodyPr vert="horz" lIns="0" tIns="0" rIns="0" bIns="0"><a:normAutofit/></a:bodyPr>'
BODY_PR_TITLE = ('<a:bodyPr vert="horz" lIns="0" tIns="0" rIns="0" bIns="0" anchor="b">'
                 "<a:normAutofit/></a:bodyPr>")
BODY_PR_CTR = '<a:bodyPr vert="horz" lIns="0" tIns="0" rIns="0" bIns="0" anchor="ctr"/>'


def ftr_sp(sid, idx=11):
    return ph_sp(
        sid, "Footer Placeholder", f'<p:ph type="ftr" sz="quarter" idx="{idx}"/>',
        ML, FTR_Y, inch(7.0), FTR_H, BODY_PR_CTR,
        lst_style_one(_sz("footer"), "accent6", bullet=False, spc_bef=0, spc=60),
        f'<a:p><a:r><a:rPr lang="ja-JP" dirty="0"/><a:t>{xml_escape(CONFIG["footer_text"])}</a:t></a:r>'
        f'<a:endParaRPr lang="ja-JP" dirty="0"/></a:p>',
    )


def sldnum_sp(sid, idx=12):
    return ph_sp(
        sid, "Slide Number Placeholder", f'<p:ph type="sldNum" sz="quarter" idx="{idx}"/>',
        SN_X, FTR_Y, SN_W, FTR_H, BODY_PR_CTR,
        lst_style_one(_sz("footer"), "accent6", bullet=False, algn="r", spc_bef=0),
        f'<a:p><a:pPr algn="r"/><a:fld id="{FLD_SLIDENUM_GUID}" type="slidenum">'
        f'<a:rPr lang="ja-JP" smtClean="0"/><a:t>2</a:t></a:fld>'
        f'<a:endParaRPr lang="ja-JP" dirty="0"/></a:p>',
    )


def hairline_sp(sid):
    return rect_sp(sid, "Footer Rule", ML, RULE_Y, CW, RULE_H, solid("accent6", 25000))


def accent_bar_sp(sid, scheme="accent1"):
    return rect_sp(sid, "Accent Bar", ML, BAR_Y, BAR_W, BAR_H, solid(scheme))


# =============================================================================
# theme1.xml
# =============================================================================

def theme_xml(name: str) -> str:
    c, f = CONFIG["colors"], CONFIG["fonts"]
    clr = "".join(
        f"<a:{k}><a:srgbClr val=\"{c[k]}\"/></a:{k}>"
        for k in ("dk1", "lt1", "dk2", "lt2", "accent1", "accent2", "accent3",
                  "accent4", "accent5", "accent6", "hlink", "folHlink")
    )
    fonts = (
        f'<a:majorFont><a:latin typeface="{f["major_latin"]}" panose="020B0502040204020203"/>'
        f'<a:ea typeface="{f["major_ea"]}"/><a:cs typeface=""/>'
        f'<a:font script="Jpan" typeface="{f["major_ea"]}"/></a:majorFont>'
        f'<a:minorFont><a:latin typeface="{f["minor_latin"]}" panose="020B0502040204020203"/>'
        f'<a:ea typeface="{f["minor_ea"]}"/><a:cs typeface=""/>'
        f'<a:font script="Jpan" typeface="{f["minor_ea"]}"/></a:minorFont>'
    )
    ln = ('<a:ln w="{w}" cap="flat" cmpd="sng" algn="ctr">'
          '<a:solidFill><a:schemeClr val="phClr"/></a:solidFill>'
          '<a:prstDash val="solid"/><a:miter lim="800000"/></a:ln>')
    fmt = (
        "<a:fillStyleLst>"
        '<a:solidFill><a:schemeClr val="phClr"/></a:solidFill>'
        '<a:solidFill><a:schemeClr val="phClr"><a:tint val="88000"/></a:schemeClr></a:solidFill>'
        '<a:solidFill><a:schemeClr val="phClr"><a:shade val="92000"/></a:schemeClr></a:solidFill>'
        "</a:fillStyleLst>"
        "<a:lnStyleLst>"
        + ln.format(w=9525) + ln.format(w=15875) + ln.format(w=25400) +
        "</a:lnStyleLst>"
        "<a:effectStyleLst>"
        "<a:effectStyle><a:effectLst/></a:effectStyle>"
        "<a:effectStyle><a:effectLst/></a:effectStyle>"
        '<a:effectStyle><a:effectLst><a:outerShdw blurRad="50800" dist="19050" dir="5400000" '
        'algn="ctr" rotWithShape="0"><a:srgbClr val="000000"><a:alpha val="10000"/></a:srgbClr>'
        "</a:outerShdw></a:effectLst></a:effectStyle>"
        "</a:effectStyleLst>"
        "<a:bgFillStyleLst>"
        '<a:solidFill><a:schemeClr val="phClr"/></a:solidFill>'
        '<a:solidFill><a:schemeClr val="phClr"><a:tint val="96000"/></a:schemeClr></a:solidFill>'
        '<a:solidFill><a:schemeClr val="phClr"/></a:solidFill>'
        "</a:bgFillStyleLst>"
    )
    return (
        XML_HEAD +
        f'<a:theme xmlns:a="{A_NS}" name="{xml_escape(name)}"><a:themeElements>'
        f'<a:clrScheme name="{xml_escape(name)}">{clr}</a:clrScheme>'
        f'<a:fontScheme name="{xml_escape(name)}">{fonts}</a:fontScheme>'
        f'<a:fmtScheme name="{xml_escape(name)}">{fmt}</a:fmtScheme>'
        f"</a:themeElements><a:objectDefaults/><a:extraClrSchemeLst/></a:theme>"
    )


# =============================================================================
# slideMaster1.xml
# =============================================================================

CLR_MAP = ('<p:clrMap bg1="lt1" tx1="dk1" bg2="lt2" tx2="dk2" accent1="accent1" '
           'accent2="accent2" accent3="accent3" accent4="accent4" accent5="accent5" '
           'accent6="accent6" hlink="hlink" folHlink="folHlink"/>')


def _tx_styles() -> str:
    title = (
        "<p:titleStyle>"
        '<a:lvl1pPr algn="l" defTabSz="914400" rtl="0" eaLnBrk="1" latinLnBrk="0" hangingPunct="1">'
        '<a:lnSpc><a:spcPct val="104000"/></a:lnSpc><a:spcBef><a:spcPct val="0"/></a:spcBef>'
        '<a:buNone/>' + def_rpr(_sz("title"), "tx1", b=1, major=True, spc=-20) +
        "</a:lvl1pPr></p:titleStyle>"
    )
    lvls = []
    for i in range(1, 10):
        if i == 1:
            lvls.append(
                '<a:lvl1pPr marL="342900" indent="-342900" algn="l" defTabSz="914400" rtl="0" '
                'eaLnBrk="1" latinLnBrk="0" hangingPunct="1">'
                '<a:lnSpc><a:spcPct val="112000"/></a:lnSpc>'
                '<a:spcBef><a:spcPts val="900"/></a:spcBef>'
                '<a:buClr><a:schemeClr val="accent1"/></a:buClr><a:buSzPct val="78000"/>'
                '<a:buFont typeface="Arial" panose="020B0604020202020204" pitchFamily="34" charset="0"/>'
                '<a:buChar char="\u25aa"/>' + def_rpr(_sz("body"), "tx2") + "</a:lvl1pPr>"
            )
        elif i == 2:
            lvls.append(
                '<a:lvl2pPr marL="742950" indent="-285750" algn="l" defTabSz="914400" rtl="0" '
                'eaLnBrk="1" latinLnBrk="0" hangingPunct="1">'
                '<a:lnSpc><a:spcPct val="112000"/></a:lnSpc>'
                '<a:spcBef><a:spcPts val="600"/></a:spcBef>'
                '<a:buClr><a:schemeClr val="accent6"/></a:buClr><a:buSzPct val="90000"/>'
                '<a:buFont typeface="Arial" panose="020B0604020202020204" pitchFamily="34" charset="0"/>'
                '<a:buChar char="\u2013"/>' + def_rpr(_sz("body") - 200, "accent6") + "</a:lvl2pPr>"
            )
        else:
            marL = 742950 + (i - 2) * 400050
            lvls.append(
                f'<a:lvl{i}pPr marL="{marL}" indent="-285750" algn="l" defTabSz="914400" rtl="0" '
                'eaLnBrk="1" latinLnBrk="0" hangingPunct="1">'
                '<a:spcBef><a:spcPts val="500"/></a:spcBef>'
                '<a:buClr><a:schemeClr val="accent6"/></a:buClr>'
                '<a:buFont typeface="Arial" pitchFamily="34" charset="0"/><a:buChar char="\u00b7"/>'
                + def_rpr(_sz("body") - 400, "accent6") + f"</a:lvl{i}pPr>"
            )
    body = "<p:bodyStyle>" + "".join(lvls) + "</p:bodyStyle>"

    other_lvls = []
    for i in range(1, 10):
        other_lvls.append(
            f'<a:lvl{i}pPr marL="{(i - 1) * 457200}" algn="l" defTabSz="914400" rtl="0" '
            'eaLnBrk="1" latinLnBrk="0" hangingPunct="1"><a:buNone/>'
            + def_rpr(1400, "tx2") + f"</a:lvl{i}pPr>"
        )
    other = ('<p:otherStyle><a:defPPr><a:defRPr lang="ja-JP"/></a:defPPr>'
             + "".join(other_lvls) + "</p:otherStyle>")
    return "<p:txStyles>" + title + body + other + "</p:txStyles>"


def master_xml() -> str:
    shapes = [
        ph_sp(2, "Title Placeholder 1", '<p:ph type="title"/>',
              ML, TITLE_Y, CW, TITLE_H, BODY_PR_TITLE),
        accent_bar_sp(3),
        ph_sp(4, "Text Placeholder 2", '<p:ph type="body" idx="1"/>',
              ML, BODY_Y, CW, BODY_H, BODY_PR_TIGHT),
        hairline_sp(5),
        ph_sp(6, "Date Placeholder 3", '<p:ph type="dt" sz="half" idx="2"/>',
              ML + inch(7.2), FTR_Y, inch(1.8), FTR_H, BODY_PR_CTR,
              lst_style_one(_sz("footer"), "accent6", bullet=False, spc_bef=0),
              f'<a:p><a:fld id="{FLD_DATE_GUID}" type="datetimeFigureOut">'
              f'<a:rPr lang="ja-JP" smtClean="0"/><a:t>2026/08/10</a:t></a:fld>'
              f'<a:endParaRPr lang="ja-JP"/></a:p>'),
        ftr_sp(7, idx=3),
        sldnum_sp(8, idx=4),
    ]
    layout_ids = "".join(
        f'<p:sldLayoutId id="{2147483649 + i}" r:id="rId{i + 1}"/>' for i in range(len(LAYOUTS))
    )
    return (
        XML_HEAD +
        f"<p:sldMaster {NS_DECL}><p:cSld>"
        f'<p:bg><p:bgPr>{solid("bg1")}<a:effectLst/></p:bgPr></p:bg>'
        f"{SPTREE_HEAD}{''.join(shapes)}{SPTREE_TAIL}</p:cSld>"
        f"{CLR_MAP}<p:sldLayoutIdLst>{layout_ids}</p:sldLayoutIdLst>"
        f"{_tx_styles()}</p:sldMaster>"
    )


# =============================================================================
# slideLayout*.xml  ―― 5レイアウト
# =============================================================================

def _layout(type_, name, shapes, bg=None) -> str:
    bg_xml = f"<p:bg><p:bgPr>{solid(bg)}<a:effectLst/></p:bgPr></p:bg>" if bg else ""
    return (
        XML_HEAD +
        f'<p:sldLayout {NS_DECL} type="{type_}" preserve="1">'
        f'<p:cSld name="{xml_escape(name)}">{bg_xml}'
        f"{SPTREE_HEAD}{''.join(shapes)}{SPTREE_TAIL}</p:cSld>"
        f"<p:clrMapOvr><a:masterClrMapping/></p:clrMapOvr></p:sldLayout>"
    )


# --- L1: 表紙 -----------------------------------------------------------------
CV_T_Y, CV_T_H = inch(1.10), inch(1.90)
CV_RULE_Y = inch(3.20)
CV_S_Y, CV_S_H = inch(3.42), inch(0.70)
CV_B_Y, CV_B_H = inch(4.30), inch(1.95)
CV_BAND_Y = inch(6.45)


def layout_cover() -> str:
    catch = xml_escape(CONFIG["catch_copy"])
    shapes = [
        rect_sp(2, "Cover Band", 0, CV_BAND_Y, SW, SH - CV_BAND_Y, solid("accent1"),
                paras=(f'<a:p><a:pPr algn="l"/><a:r><a:rPr lang="ja-JP" sz="{_sz("cover_catch")}" '
                       f'b="1" spc="300" dirty="0">{solid("bg1")}'
                       f'<a:latin typeface="+mj-lt"/><a:ea typeface="+mj-ea"/></a:rPr>'
                       f"<a:t>{catch}</a:t></a:r></a:p>"),
                body_pr=(f'<a:bodyPr lIns="{ML}" tIns="0" rIns="{ML}" bIns="0" anchor="ctr"/>')),
        ph_sp(3, "Title 1", '<p:ph type="ctrTitle"/>', ML, CV_T_Y, CW, CV_T_H,
              BODY_PR_TITLE,
              lst_style_one(_sz("cover_title"), "tx1", b=1, major=True,
                            bullet=False, spc_bef=0, ln_spc=106000, spc=-30)),
        rect_sp(4, "Cover Accent Bar", ML, CV_RULE_Y, BAR_W, BAR_H, solid("accent1")),
        ph_sp(5, "Subtitle 2", '<p:ph type="subTitle" idx="1"/>', ML, CV_S_Y, CW, CV_S_H,
              BODY_PR_TIGHT,
              lst_style_one(_sz("cover_subtitle"), "accent1", b=1, bullet=False, spc_bef=0)),
        ph_sp(6, "Text Placeholder 3", '<p:ph type="body" idx="2"/>', ML, CV_B_Y, CW, CV_B_H,
              BODY_PR_TIGHT, lst_style_one(_sz("cover_body"), "tx2", spc_bef=700)),
    ]
    return _layout("title", "表紙（タイトル）", shapes)


# --- L2: タイトルとコンテンツ ---------------------------------------------------
def layout_content() -> str:
    shapes = [
        ph_sp(2, "Title 1", '<p:ph type="title"/>', ML, TITLE_Y, CW, TITLE_H, BODY_PR_TITLE,
              lst_style_one(_sz("title"), "tx1", b=1, major=True,
                            bullet=False, spc_bef=0, ln_spc=104000, spc=-20)),
        accent_bar_sp(3),
        ph_sp(4, "Content Placeholder 2", '<p:ph type="body" idx="1"/>',
              ML, BODY_Y, CW, BODY_H, BODY_PR_TIGHT,
              lst_style_one(_sz("body"), "tx2")),
        hairline_sp(5),
        ftr_sp(6),
        sldnum_sp(7),
    ]
    return _layout("obj", "タイトルとコンテンツ", shapes)


# --- L3: 表スライド -------------------------------------------------------------
def layout_table() -> str:
    shapes = [
        ph_sp(2, "Title 1", '<p:ph type="title"/>', ML, TITLE_Y, CW, TITLE_H, BODY_PR_TITLE,
              lst_style_one(_sz("title"), "tx1", b=1, major=True,
                            bullet=False, spc_bef=0, ln_spc=104000, spc=-20)),
        accent_bar_sp(3),
        hairline_sp(4),
        ftr_sp(5),
        sldnum_sp(6),
    ]
    return _layout("tbl", "表スライド", shapes)


# --- L4: セクション見出し --------------------------------------------------------
SEC_Y = inch(2.00)
SEC_BAR_W = inch(0.09)
SEC_X = ML + inch(0.32)
SEC_W = CW - inch(0.32)


def layout_section() -> str:
    shapes = [
        rect_sp(2, "Section Bar", ML, SEC_Y, SEC_BAR_W, inch(3.00), solid("accent1")),
        ph_sp(3, "Title 1", '<p:ph type="title"/>', SEC_X, SEC_Y, SEC_W, inch(1.05),
              BODY_PR_TIGHT,
              lst_style_one(_sz("section_title"), "accent1", b=1, major=True,
                            bullet=False, spc_bef=0, ln_spc=104000, spc=-20)),
        ph_sp(4, "Text Placeholder 2", '<p:ph type="body" idx="1"/>',
              SEC_X, inch(3.20), SEC_W, inch(1.80), BODY_PR_TIGHT,
              lst_style_one(_sz("section_body"), "tx2")),
        hairline_sp(5),
        ftr_sp(6),
        sldnum_sp(7),
    ]
    return _layout("secHead", "セクション見出し", shapes)


# --- L5: クロージング（反転） ----------------------------------------------------
def layout_closing() -> str:
    shapes = [
        ph_sp(2, "Title 1", '<p:ph type="title"/>', ML, inch(1.75), CW, inch(1.00),
              BODY_PR_TITLE,
              lst_style_one(_sz("closing_title"), "bg1", b=1, major=True,
                            bullet=False, spc_bef=0, ln_spc=104000, spc=-20)),
        rect_sp(3, "Closing Rule", ML, inch(2.92), BAR_W, BAR_H, solid("bg1")),
        ph_sp(4, "Text Placeholder 2", '<p:ph type="body" idx="1"/>',
              ML, inch(3.20), CW, inch(2.60), BODY_PR_TIGHT,
              ('<a:lstStyle><a:lvl1pPr marL="342900" indent="-342900" algn="l">'
               '<a:lnSpc><a:spcPct val="112000"/></a:lnSpc>'
               '<a:spcBef><a:spcPts val="900"/></a:spcBef>'
               '<a:buClr><a:schemeClr val="bg1"><a:alpha val="70000"/></a:schemeClr></a:buClr>'
               '<a:buSzPct val="78000"/>'
               '<a:buFont typeface="Arial" panose="020B0604020202020204" pitchFamily="34" charset="0"/>'
               '<a:buChar char="▪"/>'
               + def_rpr(_sz("closing_body"), "bg1") + "</a:lvl1pPr></a:lstStyle>")),
        rect_sp(5, "Closing Rule 2", ML, RULE_Y, CW, RULE_H, solid("bg1", 35000)),
        ph_sp(6, "Footer Placeholder", '<p:ph type="ftr" sz="quarter" idx="11"/>',
              ML, FTR_Y, inch(7.0), FTR_H, BODY_PR_CTR,
              lst_style_one(_sz("footer"), "bg1", bullet=False, spc_bef=0,
                            alpha=75000, spc=60),
              f'<a:p><a:r><a:rPr lang="ja-JP" dirty="0"/>'
              f'<a:t>{xml_escape(CONFIG["footer_text"])}</a:t></a:r>'
              f'<a:endParaRPr lang="ja-JP" dirty="0"/></a:p>'),
        ph_sp(7, "Slide Number Placeholder", '<p:ph type="sldNum" sz="quarter" idx="12"/>',
              SN_X, FTR_Y, SN_W, FTR_H, BODY_PR_CTR,
              lst_style_one(_sz("footer"), "bg1", bullet=False, algn="r", spc_bef=0,
                            alpha=75000),
              f'<a:p><a:pPr algn="r"/><a:fld id="{FLD_SLIDENUM_GUID}" type="slidenum">'
              f'<a:rPr lang="ja-JP" smtClean="0"/><a:t>19</a:t></a:fld>'
              f'<a:endParaRPr lang="ja-JP" dirty="0"/></a:p>'),
    ]
    return _layout("obj", "クロージング（反転）", shapes, bg="accent1")


LAYOUTS = [layout_cover, layout_content, layout_table, layout_section, layout_closing]

L_COVER, L_CONTENT, L_TABLE, L_SECTION, L_CLOSING = 0, 1, 2, 3, 4


# =============================================================================
# notesMaster1.xml
# =============================================================================

def notes_master_xml() -> str:
    def notes_ph(sid, name, ph, x, y, cx, cy, paras="<a:p/>", extra_ln=""):
        return (
            f'<p:sp><p:nvSpPr><p:cNvPr id="{sid}" name="{name}"/>'
            f'<p:cNvSpPr><a:spLocks noGrp="1"{extra_ln}/></p:cNvSpPr>'
            f"<p:nvPr>{ph}</p:nvPr></p:nvSpPr>"
            f'<p:spPr><a:xfrm><a:off x="{x}" y="{y}"/><a:ext cx="{cx}" cy="{cy}"/></a:xfrm>'
            f'<a:prstGeom prst="rect"><a:avLst/></a:prstGeom></p:spPr>'
            f'<p:txBody><a:bodyPr vert="horz" lIns="91440" tIns="45720" rIns="91440" '
            f'bIns="45720" rtlCol="0"><a:normAutofit/></a:bodyPr><a:lstStyle/>{paras}</p:txBody></p:sp>'
        )

    img = (
        f'<p:sp><p:nvSpPr><p:cNvPr id="4" name="Slide Image Placeholder 3"/>'
        f'<p:cNvSpPr><a:spLocks noGrp="1" noRot="1" noChangeAspect="1"/></p:cNvSpPr>'
        f'<p:nvPr><p:ph type="sldImg" idx="2"/></p:nvPr></p:nvSpPr>'
        f'<p:spPr><a:xfrm><a:off x="1143000" y="685800"/><a:ext cx="4572000" cy="3429000"/></a:xfrm>'
        f'<a:prstGeom prst="rect"><a:avLst/></a:prstGeom>'
        f'<a:noFill/><a:ln w="12700"><a:solidFill><a:schemeClr val="accent6">'
        f'<a:alpha val="45000"/></a:schemeClr></a:solidFill></a:ln></p:spPr>'
        f"<p:txBody><a:bodyPr/><a:lstStyle/><a:p/></p:txBody></p:sp>"
    )
    shapes = [
        notes_ph(2, "Header Placeholder 1", '<p:ph type="hdr" sz="quarter"/>',
                 0, 0, 2971800, 458788),
        notes_ph(3, "Date Placeholder 2", '<p:ph type="dt" idx="1"/>',
                 3884613, 0, 2971800, 458788),
        img,
        notes_ph(5, "Notes Placeholder 4", '<p:ph type="body" idx="3"/>',
                 685800, 4343400, 5486400, 4114800),
        notes_ph(6, "Footer Placeholder 5", '<p:ph type="ftr" sz="quarter" idx="4"/>',
                 0, 8685213, 2971800, 458787),
        notes_ph(7, "Slide Number Placeholder 6", '<p:ph type="sldNum" sz="quarter" idx="5"/>',
                 3884613, 8685213, 2971800, 458787,
                 paras=(f'<a:p><a:pPr algn="r"/><a:fld id="{FLD_SLIDENUM_GUID}" type="slidenum">'
                        f'<a:rPr lang="ja-JP" smtClean="0"/><a:t>1</a:t></a:fld>'
                        f'<a:endParaRPr lang="ja-JP"/></a:p>')),
    ]
    lvls = "".join(
        f'<a:lvl{i}pPr marL="{(i - 1) * 457200}" algn="l" defTabSz="914400" rtl="0" '
        'eaLnBrk="1" latinLnBrk="0" hangingPunct="1">'
        '<a:lnSpc><a:spcPct val="118000"/></a:lnSpc>'
        '<a:spcBef><a:spcPts val="400"/></a:spcBef><a:buNone/>'
        + def_rpr(1200, "tx2") + f"</a:lvl{i}pPr>"
        for i in range(1, 10)
    )
    return (
        XML_HEAD +
        f"<p:notesMaster {NS_DECL}><p:cSld>"
        f'<p:bg><p:bgRef idx="1001"><a:schemeClr val="bg1"/></p:bgRef></p:bg>'
        f"{SPTREE_HEAD}{''.join(shapes)}{SPTREE_TAIL}</p:cSld>"
        f"{CLR_MAP}<p:notesStyle>{lvls}</p:notesStyle></p:notesMaster>"
    )


# =============================================================================
# パッケージ（ベーステンプレート .pptx）の組み立て
# =============================================================================

CT = {
    "pres":  "application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml",
    "sldm":  "application/vnd.openxmlformats-officedocument.presentationml.slideMaster+xml",
    "sldl":  "application/vnd.openxmlformats-officedocument.presentationml.slideLayout+xml",
    "ntm":   "application/vnd.openxmlformats-officedocument.presentationml.notesMaster+xml",
    "theme": "application/vnd.openxmlformats-officedocument.theme+xml",
    "tbls":  "application/vnd.openxmlformats-officedocument.presentationml.tableStyles+xml",
    "presp": "application/vnd.openxmlformats-officedocument.presentationml.presProps+xml",
    "viewp": "application/vnd.openxmlformats-officedocument.presentationml.viewProps+xml",
    "core":  "application/vnd.openxmlformats-package.core-properties+xml",
    "app":   "application/vnd.openxmlformats-officedocument.extended-properties+xml",
}
RT = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PKG_RT = "http://schemas.openxmlformats.org/package/2006/relationships"


def _rels(items) -> str:
    body = "".join(
        f'<Relationship Id="{rid}" Type="{typ}" Target="{tgt}"/>' for rid, typ, tgt in items
    )
    return (XML_HEAD +
            f'<Relationships xmlns="{PKG_RT}">{body}</Relationships>')


def content_types_xml() -> str:
    ov = [
        ("/ppt/presentation.xml", CT["pres"]),
        ("/ppt/slideMasters/slideMaster1.xml", CT["sldm"]),
        ("/ppt/notesMasters/notesMaster1.xml", CT["ntm"]),
        ("/ppt/theme/theme1.xml", CT["theme"]),
        ("/ppt/theme/theme2.xml", CT["theme"]),
        ("/ppt/tableStyles.xml", CT["tbls"]),
        ("/ppt/presProps.xml", CT["presp"]),
        ("/ppt/viewProps.xml", CT["viewp"]),
        ("/docProps/core.xml", CT["core"]),
        ("/docProps/app.xml", CT["app"]),
    ]
    for i in range(len(LAYOUTS)):
        ov.append((f"/ppt/slideLayouts/slideLayout{i + 1}.xml", CT["sldl"]))
    body = "".join(f'<Override PartName="{p}" ContentType="{c}"/>' for p, c in ov)
    return (
        XML_HEAD +
        '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="rels" '
        'ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
        '<Default Extension="xml" ContentType="application/xml"/>'
        f"{body}</Types>"
    )


def presentation_xml() -> str:
    lvls = "".join(
        f'<a:lvl{i}pPr marL="{(i - 1) * 457200}" algn="l" defTabSz="914400" rtl="0" '
        'eaLnBrk="1" latinLnBrk="0" hangingPunct="1">'
        + def_rpr(_sz("body"), "tx1") + f"</a:lvl{i}pPr>"
        for i in range(1, 10)
    )
    return (
        XML_HEAD +
        f'<p:presentation {NS_DECL} saveSubsetFonts="1">'
        f'<p:sldMasterIdLst><p:sldMasterId id="2147483648" r:id="rId1"/></p:sldMasterIdLst>'
        f'<p:notesMasterIdLst><p:notesMasterId r:id="rId2"/></p:notesMasterIdLst>'
        f"<p:sldIdLst/>"
        f'<p:sldSz cx="{SW}" cy="{SH}"/><p:notesSz cx="6858000" cy="9144000"/>'
        f'<p:defaultTextStyle><a:defPPr><a:defRPr lang="ja-JP"/></a:defPPr>{lvls}'
        f"</p:defaultTextStyle></p:presentation>"
    )


def core_xml() -> str:
    return (
        XML_HEAD +
        '<cp:coreProperties '
        'xmlns:cp="http://schemas.openxmlformats.org/package/2006/metadata/core-properties" '
        'xmlns:dc="http://purl.org/dc/elements/1.1/" '
        'xmlns:dcterms="http://purl.org/dc/terms/" '
        'xmlns:dcmitype="http://purl.org/dc/dcmitype/" '
        'xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">'
        f'<dc:title>{xml_escape(CONFIG["doc_title"])}</dc:title>'
        f'<dc:creator>{xml_escape(CONFIG["doc_author"])}</dc:creator>'
        f'<cp:lastModifiedBy>{xml_escape(CONFIG["doc_author"])}</cp:lastModifiedBy>'
        '<cp:revision>1</cp:revision></cp:coreProperties>'
    )


APP_XML = (
    XML_HEAD +
    '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties" '
    'xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes">'
    "<Application>Microsoft Office PowerPoint</Application>"
    "<AppVersion>16.0000</AppVersion></Properties>"
)

PRESPROPS_XML = XML_HEAD + f"<p:presentationPr {NS_DECL}/>"
VIEWPROPS_XML = XML_HEAD + f"<p:viewPr {NS_DECL}/>"
TABLESTYLES_XML = (XML_HEAD +
                   f'<a:tblStyleLst xmlns:a="{A_NS}" '
                   'def="{5C22544A-7EE6-4342-B048-85BDC9FD1C3A}"/>')


def build_base_package() -> bytes:
    """テーマ／マスター／レイアウトだけを持つ .pptx をメモリ上に組み立てる."""
    theme = theme_xml(CONFIG["theme_name"])
    parts = {
        "[Content_Types].xml": content_types_xml(),
        "_rels/.rels": _rels([
            ("rId1", f"{RT}/officeDocument", "ppt/presentation.xml"),
            ("rId2", f"{PKG_RT}/metadata/core-properties", "docProps/core.xml"),
            ("rId3", f"{RT}/extended-properties", "docProps/app.xml"),
        ]),
        "docProps/core.xml": core_xml(),
        "docProps/app.xml": APP_XML,
        "ppt/presentation.xml": presentation_xml(),
        "ppt/_rels/presentation.xml.rels": _rels([
            ("rId1", f"{RT}/slideMaster", "slideMasters/slideMaster1.xml"),
            ("rId2", f"{RT}/notesMaster", "notesMasters/notesMaster1.xml"),
            ("rId3", f"{RT}/presProps", "presProps.xml"),
            ("rId4", f"{RT}/viewProps", "viewProps.xml"),
            ("rId5", f"{RT}/theme", "theme/theme1.xml"),
            ("rId6", f"{RT}/tableStyles", "tableStyles.xml"),
        ]),
        "ppt/presProps.xml": PRESPROPS_XML,
        "ppt/viewProps.xml": VIEWPROPS_XML,
        "ppt/tableStyles.xml": TABLESTYLES_XML,
        "ppt/theme/theme1.xml": theme,
        "ppt/theme/theme2.xml": theme,
        "ppt/slideMasters/slideMaster1.xml": master_xml(),
        "ppt/slideMasters/_rels/slideMaster1.xml.rels": _rels(
            [(f"rId{i + 1}", f"{RT}/slideLayout", f"../slideLayouts/slideLayout{i + 1}.xml")
             for i in range(len(LAYOUTS))]
            + [(f"rId{len(LAYOUTS) + 1}", f"{RT}/theme", "../theme/theme1.xml")]
        ),
        "ppt/notesMasters/notesMaster1.xml": notes_master_xml(),
        "ppt/notesMasters/_rels/notesMaster1.xml.rels": _rels([
            ("rId1", f"{RT}/theme", "../theme/theme2.xml"),
        ]),
    }
    for i, builder in enumerate(LAYOUTS):
        parts[f"ppt/slideLayouts/slideLayout{i + 1}.xml"] = builder()
        parts[f"ppt/slideLayouts/_rels/slideLayout{i + 1}.xml.rels"] = _rels([
            ("rId1", f"{RT}/slideMaster", "../slideMasters/slideMaster1.xml"),
        ])

    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for name, xml in parts.items():
            z.writestr(name, xml.encode("utf-8"))
    return buf.getvalue()


def open_base() -> Presentation:
    return Presentation(io.BytesIO(build_base_package()))


# =============================================================================
# Markdown パース
# =============================================================================

NOTE_RE = re.compile(r"<!--\s*note:\s*(.*?)\s*-->", re.S)
BOLD_RE = re.compile(r"(\*\*.+?\*\*)", re.S)


class SlideSpec:
    def __init__(self):
        self.title = ""
        self.subtitle = ""
        self.items = []      # [(level:int, text:str)]
        self.paras = []      # 箇条書き以外の段落
        self.table = None    # [[cell, ...], ...]  1行目がヘッダー
        self.note = ""


def parse_markdown(path: str):
    with open(path, encoding="utf-8") as fh:
        text = fh.read()

    blocks = re.split(r"^---\s*$", text, flags=re.M)
    specs = []
    for block in blocks:
        if not block.strip():
            continue
        spec = SlideSpec()
        m = NOTE_RE.search(block)
        if m:
            spec.note = re.sub(r"\s*\n\s*", " ", m.group(1)).strip()
            block = NOTE_RE.sub("", block)

        table_rows = []
        for raw in block.splitlines():
            line = raw.rstrip()
            if not line.strip():
                continue
            if line.startswith("## "):
                spec.subtitle = line[3:].strip()
            elif line.startswith("# "):
                spec.title = line[2:].strip()
            elif line.lstrip().startswith("|"):
                cells = [c.strip() for c in line.strip().strip("|").split("|")]
                if all(set(c) <= set("-: ") and c for c in cells):
                    continue                      # 区切り行
                table_rows.append(cells)
            else:
                stripped = line.lstrip(" ")
                indent = len(line) - len(stripped)
                if stripped.startswith("- "):
                    spec.items.append((1 if indent >= 2 else 0, stripped[2:].strip()))
                else:
                    spec.paras.append(line.strip())
        if table_rows:
            spec.table = table_rows
        specs.append(spec)
    return specs


# =============================================================================
# テキスト描画ヘルパー
# =============================================================================

def _set_theme_font(rPr, major=False):
    pre = "mj" if major else "mn"
    for tag, val in (("a:latin", f"+{pre}-lt"), ("a:ea", f"+{pre}-ea"), ("a:cs", f"+{pre}-cs")):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            rPr.append(el)
        el.set("typeface", val)


def add_runs(paragraph, text, size_pt=None, major=False, bold=False):
    """`**太字**` を解釈して run を追加する."""
    pieces = [p for p in BOLD_RE.split(text) if p != ""]
    if not pieces:
        pieces = [""]
    for piece in pieces:
        is_bold = bold
        if piece.startswith("**") and piece.endswith("**") and len(piece) > 4:
            piece = piece[2:-2]
            is_bold = True
        run = paragraph.add_run()
        run.text = piece
        if size_pt is not None:
            run.font.size = Pt(size_pt)
        if is_bold:
            run.font.bold = True
        rPr = run._r.get_or_add_rPr()
        rPr.set("lang", "ja-JP")
        rPr.set("dirty", "0")
        _set_theme_font(rPr, major)


NUMBERED_RE = re.compile(r"^\(\d+\)")


def _no_bullet(paragraph, marL=342900, indent=-342900):
    pPr = paragraph._p.get_or_add_pPr()
    pPr.set("marL", str(marL))
    pPr.set("indent", str(indent))
    for tag in ("a:buNone", "a:buChar", "a:buAutoNum"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    bu = pPr.makeelement(qn("a:buNone"), {})
    pPr.insert_element_before(bu, "a:tabLst", "a:defRPr", "a:extLst")


def fill_body(placeholder, items, paras, size_pt):
    tf = placeholder.text_frame
    tf.clear()
    first = True
    entries = [(0, p, False) for p in paras] + [(lv, t, True) for lv, t in items]
    for level, text, is_bullet in entries:
        para = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        para.level = level
        if not is_bullet:
            _no_bullet(para, marL=0, indent=0)
        elif NUMBERED_RE.match(text):
            # `- (1) 項目` は自動番号にせず、そのままの文字として字下げ表示
            _no_bullet(para, marL=342900, indent=-342900)
        add_runs(para, text, size_pt=size_pt)


def fill_title(slide, text, size_pt=None):
    ph = slide.shapes.title
    tf = ph.text_frame
    tf.clear()
    add_runs(tf.paragraphs[0], text, size_pt=size_pt, major=True)
    return ph


def set_notes(slide, text):
    tf = slide.notes_slide.notes_text_frame
    tf.clear()
    add_runs(tf.paragraphs[0], text, size_pt=12)


def copy_footer_shapes(slide, layout):
    """レイアウトのフッター／スライド番号プレースホルダーをスライドへ複製する."""
    sp_tree = slide.shapes._spTree
    ids = [int(e.get("id")) for e in sp_tree.iter(qn("p:cNvPr")) if e.get("id")]
    next_id = max(ids or [1]) + 1
    for sp in layout.shapes._spTree:
        if sp.tag != qn("p:sp"):
            continue
        ph = sp.find(qn("p:nvSpPr") + "/" + qn("p:nvPr") + "/" + qn("p:ph"))
        if ph is None or ph.get("type") not in ("ftr", "sldNum"):
            continue
        new = copy.deepcopy(sp)
        cNvPr = new.find(qn("p:nvSpPr")).find(qn("p:cNvPr"))
        cNvPr.set("id", str(next_id))
        next_id += 1
        for fld in new.iter(qn("a:fld")):
            fld.set("id", "{%s}" % str(uuid.uuid4()).upper())
        sp_tree.append(new)


# =============================================================================
# 表の描画
# =============================================================================

FILL_TAGS = ("a:noFill", "a:solidFill", "a:gradFill", "a:blipFill", "a:pattFill", "a:grpFill")


def cell_fill(cell, scheme, alpha=None):
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in FILL_TAGS:
        for el in tcPr.findall(qn(tag)):
            tcPr.remove(el)
    inner = f'<a:alpha val="{alpha}"/>' if alpha is not None else ""
    xml = (f'<a:solidFill {nsdecls("a")}><a:schemeClr val="{scheme}">{inner}'
           f"</a:schemeClr></a:solidFill>")
    tcPr.insert_element_before(parse_xml(xml), "a:headers", "a:extLst")


def cell_borders(cell, top=None, bottom=None, w=9525):
    """top/bottom は (schemeClr, alpha) または None."""
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        for el in tcPr.findall(qn(tag)):
            tcPr.remove(el)

    def ln(tag, spec):
        if spec is None:
            fill = "<a:noFill/>"
        else:
            scheme, alpha = spec
            inner = f'<a:alpha val="{alpha}"/>' if alpha is not None else ""
            fill = (f'<a:solidFill><a:schemeClr val="{scheme}">{inner}'
                    f"</a:schemeClr></a:solidFill>")
        return parse_xml(
            f'<{tag} {nsdecls("a")} w="{w}" cap="flat" cmpd="sng" algn="ctr">{fill}'
            f'<a:prstDash val="solid"/><a:round/>'
            f'<a:headEnd type="none" w="med" len="med"/>'
            f'<a:tailEnd type="none" w="med" len="med"/></{tag}>'
        )

    # 挿入順を L, R, T, B にするため逆順に先頭挿入
    tcPr.insert(0, ln("a:lnB", bottom))
    tcPr.insert(0, ln("a:lnT", top))
    tcPr.insert(0, ln("a:lnR", None))
    tcPr.insert(0, ln("a:lnL", None))


def add_table(slide, rows_data):
    n_rows, n_cols = len(rows_data), len(rows_data[0])
    head_h, body_h = inch(0.46), inch(0.42)
    total_h = head_h + body_h * (n_rows - 1)
    gf = slide.shapes.add_table(n_rows, n_cols, Emu(ML), Emu(TBL_Y), Emu(CW), Emu(total_h))
    table = gf.table
    table.first_row = False
    table.horz_banding = False

    if n_cols == 2:
        table.columns[0].width = Emu(TBL_COL0)
        table.columns[1].width = Emu(TBL_COL1)
    else:
        for col in table.columns:
            col.width = Emu(CW // n_cols)

    rule = ("accent6", 25000)
    for r, row_data in enumerate(rows_data):
        table.rows[r].height = Emu(head_h if r == 0 else body_h)
        for c, value in enumerate(row_data):
            cell = table.cell(r, c)
            cell.margin_left = Emu(inch(0.16))
            cell.margin_right = Emu(inch(0.16))
            cell.margin_top = Emu(inch(0.04))
            cell.margin_bottom = Emu(inch(0.04))
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE

            if r == 0:
                cell_fill(cell, "accent1")
                cell_borders(cell, top=("accent1", None), bottom=("accent1", None))
            elif r % 2 == 0:
                cell_fill(cell, "accent6", alpha=6000)
                cell_borders(cell, top=rule, bottom=rule)
            else:
                cell_fill(cell, "bg1")
                cell_borders(cell, top=rule, bottom=rule)

            tf = cell.text_frame
            tf.clear()
            para = tf.paragraphs[0]
            _no_bullet(para, marL=0, indent=0)
            if r == 0:
                add_runs(para, value, size_pt=CONFIG["sizes"]["table_header"], bold=True)
                for run in para.runs:
                    run.font.color.theme_color = _theme_color("bg1")
            else:
                add_runs(para, value, size_pt=CONFIG["sizes"]["table_body"],
                         bold=(c == 0))
                for run in para.runs:
                    run.font.color.theme_color = _theme_color(
                        "accent1" if c == 0 else "tx2")
    return gf


def _theme_color(name):
    from pptx.enum.dml import MSO_THEME_COLOR
    return {
        "bg1": MSO_THEME_COLOR.BACKGROUND_1,
        "tx1": MSO_THEME_COLOR.TEXT_1,
        "tx2": MSO_THEME_COLOR.TEXT_2,
        "accent1": MSO_THEME_COLOR.ACCENT_1,
        "accent6": MSO_THEME_COLOR.ACCENT_6,
    }[name]


# =============================================================================
# デッキ生成
# =============================================================================

# md の並び（0 始まり）に対するレイアウト割り当て
LAYOUT_FOR_SLIDE = {
    0:  L_COVER,     # 表紙
    2:  L_TABLE,     # 本日のアジェンダ
    8:  L_TABLE,     # Day1
    9:  L_TABLE,     # Day2
    17: L_SECTION,   # Q&A
    18: L_CLOSING,   # クロージング
}


def build_deck(specs, path):
    prs = open_base()
    prs.core_properties.title = CONFIG["doc_title"]
    prs.core_properties.author = CONFIG["doc_author"]
    sizes = CONFIG["sizes"]

    for idx, spec in enumerate(specs):
        li = LAYOUT_FOR_SLIDE.get(idx, L_CONTENT)
        layout = prs.slide_layouts[li]
        slide = prs.slides.add_slide(layout)

        if li == L_COVER:
            fill_title(slide, spec.title)
            slide.placeholders[1].text_frame.clear()
            add_runs(slide.placeholders[1].text_frame.paragraphs[0], spec.subtitle,
                     size_pt=sizes["cover_subtitle"])
            fill_body(slide.placeholders[2], spec.items, spec.paras, sizes["cover_body"])
        elif li == L_TABLE:
            fill_title(slide, spec.title)
            if spec.table:
                add_table(slide, spec.table)
        elif li == L_SECTION:
            fill_title(slide, spec.title)
            fill_body(slide.placeholders[1], spec.items, spec.paras, sizes["section_body"])
        elif li == L_CLOSING:
            fill_title(slide, spec.title)
            fill_body(slide.placeholders[1], spec.items, spec.paras, sizes["closing_body"])
        else:
            fill_title(slide, spec.title)
            n = len(spec.items) + len(spec.paras)
            size = sizes["body_large"] if n <= 4 else sizes["body"]
            fill_body(slide.placeholders[1], spec.items, spec.paras, size)

        copy_footer_shapes(slide, layout)
        if spec.note:
            set_notes(slide, spec.note)

    prs.save(path)
    return path


TEMPLATE_SAMPLE = {
    "title": "資料タイトルをここに入力",
    "subtitle": "サブタイトル／開催日・場所など",
    "items": [
        "このファイルは『ことなら』ブランドのテンプレートです",
        "［ホーム］＞［新しいスライド］から5種類のレイアウトを選べます",
        "配色・フォントは［デザイン］＞［バリエーション］で差し替えできます",
    ],
}


def build_template(path):
    prs = open_base()
    prs.core_properties.title = "ことならテック ブランドテンプレート"
    prs.core_properties.author = CONFIG["doc_author"]
    layout = prs.slide_layouts[L_COVER]
    slide = prs.slides.add_slide(layout)
    fill_title(slide, TEMPLATE_SAMPLE["title"])
    tf = slide.placeholders[1].text_frame
    tf.clear()
    add_runs(tf.paragraphs[0], TEMPLATE_SAMPLE["subtitle"],
             size_pt=CONFIG["sizes"]["cover_subtitle"])
    fill_body(slide.placeholders[2], [(0, t) for t in TEMPLATE_SAMPLE["items"]], [],
              CONFIG["sizes"]["cover_body"])
    prs.save(path)
    return path


# =============================================================================
# 検証
# =============================================================================

def _all_text(prs):
    chunks = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                chunks.append(shape.text_frame.text)
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        chunks.append(cell.text_frame.text)
    return "\n".join(chunks)


def verify(deck_path, template_path, specs):
    ok = True
    results = []

    def check(label, cond, detail=""):
        nonlocal ok
        ok = ok and bool(cond)
        results.append(f"  [{'OK' if cond else 'NG'}] {label}{(' — ' + detail) if detail else ''}")

    for p in (deck_path, template_path):
        with zipfile.ZipFile(p) as z:
            check(f"zip 健全性: {os.path.basename(p)}", z.testzip() is None)
            names = z.namelist()
            check(f"パート重複なし: {os.path.basename(p)}", len(names) == len(set(names)))
            theme = z.read("ppt/theme/theme1.xml").decode("utf-8")
            c = CONFIG["colors"]
            for key in ("accent1", "accent2", "accent3", "accent4", "accent5",
                        "accent6", "dk1", "lt1", "dk2", "lt2", "hlink", "folHlink"):
                check(f"theme1.xml clrScheme {key}={c[key]}",
                      f"<a:{key}><a:srgbClr val=\"{c[key]}\"/></a:{key}>" in theme)
            f = CONFIG["fonts"]
            check("theme1.xml majorFont latin", f'<a:latin typeface="{f["major_latin"]}"' in theme)
            check("theme1.xml majorFont ea", f'<a:ea typeface="{f["major_ea"]}"/>' in theme)
            check("theme1.xml minorFont latin", f'<a:latin typeface="{f["minor_latin"]}"' in theme)
            check("theme1.xml fmtScheme", "<a:fmtScheme" in theme)
            # 参照切れチェック
            for rel_name in [n for n in names if n.endswith(".rels")]:
                base = os.path.dirname(os.path.dirname(rel_name))
                body = z.read(rel_name).decode("utf-8")
                for tgt in re.findall(r'Target="([^"]+)"', body):
                    if tgt.startswith("http") or "TargetMode" in tgt:
                        continue
                    resolved = os.path.normpath(os.path.join(base, tgt)).replace("\\", "/")
                    resolved = resolved.lstrip("/")
                    if resolved not in names:
                        check(f"参照切れ {rel_name} -> {tgt}", False)

    prs = Presentation(deck_path)
    check("本編スライド数 = 19", len(prs.slides) == 19, f"実際={len(prs.slides)}")
    check("md スライド数 = 19", len(specs) == 19, f"実際={len(specs)}")

    no_notes = [i + 1 for i, s in enumerate(prs.slides)
                if not (s.has_notes_slide and s.notes_slide.notes_text_frame.text.strip())]
    check("全スライドにノートあり", not no_notes, f"欠落={no_notes}")

    tables = sum(1 for s in prs.slides for sh in s.shapes
                 if getattr(sh, "has_table", False) and sh.has_table)
    check("ネイティブ表 3枚", tables == 3, f"実際={tables}")

    text = _all_text(prs)
    check("`**` 残骸なし", "**" not in text)
    check("`<!-- note:` 残骸なし", "<!-- note:" not in text and "-->" not in text)

    md_titles = [s.title for s in specs]
    pptx_titles = [s.shapes.title.text_frame.text for s in prs.slides]
    mismatch = [(i + 1, a, b) for i, (a, b) in enumerate(zip(md_titles, pptx_titles)) if a != b]
    check("タイトル19本が md と一致", not mismatch, str(mismatch[:3]))

    layouts_used = sorted({s.slide_layout.name for s in prs.slides})
    check("レイアウト継承あり（手置きテキストボックス化していない）",
          all(any(sh.is_placeholder for sh in s.shapes) for s in prs.slides))
    results.append(f"  ... 使用レイアウト: {', '.join(layouts_used)}")

    tprs = Presentation(template_path)
    check("テンプレート スライド数 = 1", len(tprs.slides) == 1)
    check("テンプレート レイアウト数 = 5", len(tprs.slide_layouts) == len(LAYOUTS))
    check("テンプレート マスター数 = 1", len(tprs.slide_masters) == 1)

    print("\n".join(results))
    print("\n検証結果: " + ("ALL OK" if ok else "*** 失敗あり ***"))
    return ok


# =============================================================================
def main():
    specs = parse_markdown(MD_PATH)
    print(f"md から {len(specs)} スライドを読み込みました: {MD_PATH}")
    build_deck(specs, DECK_PATH)
    print(f"生成: {DECK_PATH}")
    build_template(TEMPLATE_PATH)
    print(f"生成: {TEMPLATE_PATH}")
    print("-" * 60)
    return 0 if verify(DECK_PATH, TEMPLATE_PATH, specs) else 1


if __name__ == "__main__":
    raise SystemExit(main())
