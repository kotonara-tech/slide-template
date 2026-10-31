# -*- coding: utf-8 -*-
"""
build_kotonara_theme_pptx.py
============================
ことなら（kotonara.tech）ブランドテーマで、Markdown 原稿から PPTX を作る。
イベント資料・営業資料などに使う汎用テンプレート。フッターの文言は既定で空欄。

雛形は3種類（リポジトリ直下）:
  kotonara-brand-template.pptx          メイン。メインの緑 #4D9622、「ことテク」ロゴ入り
  kotonara-brand-template-nologo-1.pptx ロゴなし版1。メインと同じテーマでロゴなし
  kotonara-brand-template-nologo-2.pptx ロゴなし版2。旧テンプレート（緑 #1A832A）。このスクリプトでは作らない

ビルド機構は hackathon-2026/scripts/build_hackathon_slide_pptx.py を読み込んで使う
（テーマ・5レイアウト・md の書式・表の描画）。読み込むだけで、そのファイルには書き込まない。
このスクリプトで上書きするのは次の点だけ:

  - メインの緑（accent1）: 4D9622
  - フッター: 既定は空欄（--footer で指定）
  - 表紙の帯のキャッチコピー「ひらめきを、カタチに。」: 13pt
  - 「ことテク」ロゴをスライドレイアウトに置く（--no-logo で置かない）
      表紙: 帯の左に白いロゴ。キャッチコピーはロゴの右に置き、下端をロゴの下端にそろえる
      タイトルとコンテンツ／表スライド／セクション見出し: 右上に緑のロゴ
      クロージング（反転）: 置かない
    PNG を表示し、SVG を高画質版として埋め込む（PowerPoint 2016 以降は SVG で描画する）
  - レイアウトの割り当て: 1枚目は表紙、最後はクロージング、表があれば表スライド、
    それ以外はタイトルとコンテンツ。`<!-- layout: section -->` を書いた枚はセクション見出し
  - カレンダー表: ヘッダーが「日 | 月 | 火 | 水 | 木 | 金 | 土 | この週のこと」の表は
    月のカレンダーとして描く（後述）
  - 表は表スライドにだけ置ける。表スライドには箇条書き・本文を書けない。
    どちらも内容が消えるのを防ぐため、違反するとエラーで止める

使い方（リポジトリルートで）:
  python scripts/build_kotonara_theme_pptx.py <md> [-o out.pptx] [--footer 文言] [--no-logo] [--draft]

  -o を省くと <md の場所>/<md 名>-kotonara.pptx に出力する。
  --draft を付けると {{...}} の残存を許す（社内確認用）。付けないと、1つでも残っていれば失敗する。

雛形2種（メインとロゴなし版1）の再生成:
  python scripts/build_kotonara_theme_pptx.py --build-templates

カレンダー表の書き方:
  | 日 | 月 | 火 | 水 | 木 | 金 | 土 | この週のこと |
  | --- | --- | --- | --- | --- | --- | --- | --- |
  |  |  | 1 | 2 | 3 | 4 | 5 |  |
  | 6 | 7 | 8 | 9 | 10 | 11 | 12 | 資料のご提出（11日まで） |
  | 13 | 14 | 15 | 16 | 17 | 18 | 19 当日 | 当日：11:00 ご入場 |
  - 日付セルは「19」または「19 ラベル」（空白は全角でもよい）。ラベル付きはベージュ、
    ラベルに「当日」を含むと緑で塗る
  - 列はヘッダーと同じ8列。ヘッダーのほかに1行以上の週が要る
  - 空セルは月の外（灰色）
  - 「この週のこと」が縦に同じ文言で続くと、セルを結合する
"""

from __future__ import annotations

import argparse
import importlib.util
import os
import re
import sys
import tempfile
import zipfile

from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.opc.package import Part
from pptx.opc.packuri import PackURI
from pptx.util import Emu

REPO_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
BASE_SCRIPT = os.path.join(REPO_ROOT, "hackathon-2026", "scripts", "build_hackathon_slide_pptx.py")
ASSETS = os.path.join(REPO_ROOT, "assets")
TEMPLATE_MD = os.path.join(REPO_ROOT, "kotonara-brand-template.md")
TEMPLATE_MAIN = os.path.join(REPO_ROOT, "kotonara-brand-template.pptx")
TEMPLATE_NOLOGO_1 = os.path.join(REPO_ROOT, "kotonara-brand-template-nologo-1.pptx")

ACCENT1 = "4D9622"
COVER_CATCH_PT = 13
LOGO = (os.path.join(ASSETS, "kotonara-logo.png"), os.path.join(ASSETS, "kotonara-logo.svg"))
LOGO_WHITE = (os.path.join(ASSETS, "kotonara-logo-white.png"),
              os.path.join(ASSETS, "kotonara-logo-white.svg"))

SVG_EXT_URI = "{96DAC541-7B7A-43D3-8B79-37D633B846F1}"
ASVG_NS = "http://schemas.microsoft.com/office/drawing/2016/SVG/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"

LAYOUT_RE = re.compile(r"<!--\s*layout:\s*(\w+)\s*-->")
WEEKDAYS = ["日", "月", "火", "水", "木", "金", "土"]
DAY_LABEL_RE = re.compile(r"[\s\u3000]+")   # 日付とラベルの区切り（全角空白も可）


def load_base(footer: str):
    spec = importlib.util.spec_from_file_location("kotonara_base", BASE_SCRIPT)
    B = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(B)
    B.CONFIG["colors"]["accent1"] = ACCENT1
    B.CONFIG["sizes"]["cover_catch"] = COVER_CATCH_PT
    B.CONFIG["footer_text"] = footer
    B.CONFIG["doc_title"] = "ことならテック"
    B.CONFIG["doc_author"] = "ことならテックプロジェクト"
    return B


# =============================================================================
# ロゴ
# =============================================================================

class Logo:
    """ロゴの PNG（表示用）と SVG（高画質版）。SVG のパートはデッキ内で1つを共有する。"""

    def __init__(self, png, svg):
        from PIL import Image

        self.png, self.svg = png, svg
        self._svg_part = None
        with Image.open(png) as im:
            w, h = im.size
        self.ratio = w / h

    def add_to_layout(self, prs, layout, x, y, height, name="ことテク ロゴ"):
        from pptx.oxml.shapes.picture import CT_Picture

        width = int(height * self.ratio)
        _part, rid_png = layout.part.get_or_add_image_part(self.png)
        tree = layout.shapes._spTree
        ids = [int(e.get("id")) for e in tree.iter() if e.tag.endswith("}cNvPr") and e.get("id")]
        pic = CT_Picture.new_pic(max(ids or [1]) + 1, name, "ことテク", rid_png,
                                 Emu(x), Emu(y), Emu(width), Emu(height))
        tree.append(pic)   # レイアウト上の帯などより前面
        if self.svg:
            if self._svg_part is None:
                with open(self.svg, "rb") as fh:
                    self._svg_part = Part(PackURI("/ppt/media/" + os.path.basename(self.svg)),
                                          "image/svg+xml", prs.part.package, fh.read())
            rid_svg = layout.part.relate_to(self._svg_part, RT.IMAGE)
            blip = pic.find(f".//{{{A_NS}}}blip")
            ext_lst = etree.SubElement(blip, f"{{{A_NS}}}extLst")
            ext = etree.SubElement(ext_lst, f"{{{A_NS}}}ext", uri=SVG_EXT_URI)
            svg_blip = etree.SubElement(ext, f"{{{ASVG_NS}}}svgBlip", nsmap={"asvg": ASVG_NS})
            svg_blip.set(f"{{{R_NS}}}embed", rid_svg)
        return pic


def place_logos(B, prs):
    logo, white = Logo(*LOGO), Logo(*LOGO_WHITE)

    # 表紙: 帯の左に白いロゴ。キャッチコピーはロゴの右、下端をロゴの下端にそろえる
    band_h = B.SH - B.CV_BAND_Y
    cover_h = int(band_h * 0.62)
    cover_w = int(cover_h * white.ratio)
    gap = B.inch(0.12)
    margin = (band_h - cover_h) // 2
    cover_layout = prs.slide_layouts[B.L_COVER]
    for sp in cover_layout.shapes:
        if sp.name == "Cover Band":
            body_pr = sp.text_frame._txBody.bodyPr
            body_pr.set("lIns", str(B.ML + cover_w + gap))
            body_pr.set("anchor", "b")
            body_pr.set("bIns", str(margin))
    white.add_to_layout(prs, cover_layout, B.ML, B.CV_BAND_Y + margin, cover_h)

    # 本文・表・セクション見出し: 右上。右端はコンテンツ幅の右端にそろえる
    top_h = B.inch(0.50)
    top_w = int(top_h * logo.ratio)
    for li in (B.L_CONTENT, B.L_TABLE, B.L_SECTION):
        logo.add_to_layout(prs, prs.slide_layouts[li], B.ML + B.CW - top_w, B.inch(0.40), top_h)

    # 本文・表はタイトル枠が全幅でロゴと重なるので、ロゴの左で折り返すよう幅を縮める
    # （セクション見出しのタイトルはロゴより下にあるので縮めない）
    for li in (B.L_CONTENT, B.L_TABLE):
        for sp in prs.slide_layouts[li].placeholders:
            if sp.placeholder_format.type == PP_PLACEHOLDER.TITLE:
                sp.width = Emu(B.CW - top_w - gap)


# =============================================================================
# カレンダー表
# =============================================================================

def is_calendar(table) -> bool:
    """ヘッダーが曜日7列で始まる表をカレンダーとみなす。形が崩れていればエラーで止める。"""
    if not table or table[0][:7] != WEEKDAYS:
        return False
    n_cols = len(WEEKDAYS) + 1
    if len(table[0]) != n_cols:
        raise SystemExit(f"カレンダー表のヘッダーは「{' | '.join(WEEKDAYS)} | この週のこと」"
                         f"の{n_cols}列にしてください（実際={len(table[0])}列）")
    if len(table) < 2:
        raise SystemExit("カレンダー表にヘッダーのほか1行以上の週を書いてください")
    bad = [r for r, row in enumerate(table) if len(row) != n_cols]
    if bad:
        raise SystemExit(f"カレンダー表の {', '.join(str(r + 1) for r in bad)} 行目の列数が"
                         f"ヘッダーの{n_cols}列と合いません")
    return True


def add_calendar(B, slide, rows):
    n_rows, n_cols = len(rows), len(rows[0])
    day_w = B.inch(1.12)
    week_w = B.CW - day_w * 7
    head_h = B.inch(0.42)
    body_h = (B.RULE_Y - B.inch(0.10) - B.TBL_Y - head_h) // (n_rows - 1)
    gf = slide.shapes.add_table(n_rows, n_cols, Emu(B.ML), Emu(B.TBL_Y),
                                Emu(B.CW), Emu(head_h + body_h * (n_rows - 1)))
    table = gf.table
    table.first_row = False
    table.horz_banding = False
    for c in range(7):
        table.columns[c].width = Emu(day_w)
    table.columns[7].width = Emu(week_w)

    rule = ("accent6", 25000)
    for r, row in enumerate(rows):
        table.rows[r].height = Emu(head_h if r == 0 else body_h)
        for c, value in enumerate(row):
            cell = table.cell(r, c)
            cell.margin_left = Emu(B.inch(0.05))
            cell.margin_right = Emu(B.inch(0.05))
            cell.margin_top = Emu(B.inch(0.05))
            cell.margin_bottom = Emu(B.inch(0.04))
            tf = cell.text_frame
            tf.clear()
            tf.word_wrap = True
            p = tf.paragraphs[0]
            B._no_bullet(p, marL=0, indent=0)

            if r == 0:
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                B.cell_fill(cell, "accent1")
                B.cell_borders(cell, top=("accent1", None), bottom=("accent1", None))
                p.alignment = PP_ALIGN.CENTER if c < 7 else PP_ALIGN.LEFT
                B.add_runs(p, value, size_pt=13, bold=True)
                for run in p.runs:
                    run.font.color.theme_color = B._theme_color("bg1")
                continue

            cell.vertical_anchor = MSO_ANCHOR.TOP
            B.cell_borders(cell, top=rule, bottom=rule)
            if c == 7:
                B.cell_fill(cell, "bg1")
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                B.add_runs(p, value, size_pt=12)
                for run in p.runs:
                    run.font.color.theme_color = B._theme_color("tx2")
                continue

            if not value:
                B.cell_fill(cell, "accent6", alpha=6000)   # 月の外
                continue

            day, label = (DAY_LABEL_RE.split(value.strip(), maxsplit=1) + [""])[:2]
            is_event_day = "当日" in label
            if is_event_day:
                B.cell_fill(cell, "accent1")
            elif label:
                B.cell_fill(cell, "bg2")
            else:
                B.cell_fill(cell, "bg1")
            B.add_runs(p, day, size_pt=14, bold=True, major=True)
            num_color = "bg1" if is_event_day else ("accent1" if label else "tx2")
            for run in p.runs:
                run.font.color.theme_color = B._theme_color(num_color)
            if label:
                p2 = tf.add_paragraph()
                B._no_bullet(p2, marL=0, indent=0)
                B.add_runs(p2, label, size_pt=10, bold=True)
                for run in p2.runs:
                    run.font.color.theme_color = B._theme_color("bg1" if is_event_day else "tx1")

    # 「この週のこと」: 同じ文言が縦に続くセルを結合する
    r = 1
    while r < n_rows:
        text = rows[r][7]
        end = r
        while text and end + 1 < n_rows and rows[end + 1][7] == text:
            end += 1
        if end > r:
            for k in range(r + 1, end + 1):
                tf = table.cell(k, 7).text_frame
                tf.clear()
                B._no_bullet(tf.paragraphs[0], marL=0, indent=0)
            table.cell(r, 7).merge(table.cell(end, 7))
        r = end + 1
    return gf


# =============================================================================
# md の読み込みとデッキ生成
# =============================================================================

def parse(B, md_path):
    """元スクリプトの parse_markdown に、`<!-- layout: ... -->` の指定を足す。

    指定はスライドごとに取り出して消してから元の parse_markdown に渡す。
    指定だけで中身のない枚はエラーにする（数え方がずれて別の枚に効くのを防ぐ）。
    元の parse_markdown はパスしか受け取らず、元ファイルは書き換えないので一時ファイルを使う。
    """
    with open(md_path, encoding="utf-8") as fh:
        text = fh.read()
    blocks, layouts = [], []
    for b in re.split(r"^---\s*$", text, flags=re.M):
        if not b.strip():
            continue
        m = LAYOUT_RE.search(b)
        b = LAYOUT_RE.sub("", b)
        if not b.strip():
            raise SystemExit(f"{len(blocks) + 1} 枚目: レイアウト指定だけで中身がありません")
        blocks.append(b)
        layouts.append(m.group(1) if m else None)
    with tempfile.NamedTemporaryFile("w", suffix=".md", delete=False, encoding="utf-8") as tmp:
        tmp.write("\n---\n".join(blocks))
    try:
        specs = B.parse_markdown(tmp.name)
    finally:
        os.remove(tmp.name)
    if len(specs) != len(layouts):
        raise SystemExit(f"スライド区切りの数が合いません（md={len(layouts)}, 読み込み={len(specs)}）")
    return specs, layouts


def layout_index(B, idx, last, spec, override):
    names = {"cover": B.L_COVER, "content": B.L_CONTENT, "table": B.L_TABLE,
             "section": B.L_SECTION, "closing": B.L_CLOSING}
    if override:
        if override not in names:
            raise SystemExit(f"不明なレイアウト指定: {override}（{', '.join(names)}）")
        return names[override]
    if idx == 0:
        return B.L_COVER
    if idx == last:
        return B.L_CLOSING
    if spec.table:
        return B.L_TABLE
    return B.L_CONTENT


def check_content_fits(B, idx, spec, li):
    """レイアウトに入らない内容があれば、黙って捨てずにエラーで止める。"""
    where = f"{idx + 1} 枚目「{spec.title}」"
    if spec.table and li != B.L_TABLE:
        raise SystemExit(f"{where}: 表は表スライドにしか置けません。表を別の枚に分けるか、"
                         "表紙・最後の枚・レイアウト指定のある枚から表を外してください")
    if li == B.L_TABLE and (spec.items or spec.paras):
        raise SystemExit(f"{where}: 表スライドには箇条書き・本文を書けません。"
                         "別の枚に分けるか、ノート（<!-- note: ... -->）に移してください")


def build_deck(B, specs, layouts, path, with_logo=True, doc_title=None):
    prs = B.open_base()
    if with_logo:
        place_logos(B, prs)
    prs.core_properties.title = doc_title or " ".join(
        t for t in (specs[0].title, specs[0].subtitle) if t)
    prs.core_properties.author = B.CONFIG["doc_author"]
    sizes = B.CONFIG["sizes"]
    last = len(specs) - 1

    for idx, spec in enumerate(specs):
        li = layout_index(B, idx, last, spec, layouts[idx])
        check_content_fits(B, idx, spec, li)
        layout = prs.slide_layouts[li]
        slide = prs.slides.add_slide(layout)
        B.fill_title(slide, spec.title)

        if li == B.L_COVER:
            slide.placeholders[1].text_frame.clear()
            B.add_runs(slide.placeholders[1].text_frame.paragraphs[0], spec.subtitle,
                       size_pt=sizes["cover_subtitle"])
            B.fill_body(slide.placeholders[2], spec.items, spec.paras, sizes["cover_body"])
        elif li == B.L_TABLE:
            if spec.table:
                if is_calendar(spec.table):
                    add_calendar(B, slide, spec.table)
                else:
                    B.add_table(slide, spec.table)
        elif li == B.L_SECTION:
            B.fill_body(slide.placeholders[1], spec.items, spec.paras, sizes["section_body"])
        elif li == B.L_CLOSING:
            B.fill_body(slide.placeholders[1], spec.items, spec.paras, sizes["closing_body"])
        else:
            n = len(spec.items) + len(spec.paras)
            B.fill_body(slide.placeholders[1], spec.items, spec.paras,
                        sizes["body_large"] if n <= 4 else sizes["body"])

        B.copy_footer_shapes(slide, layout)
        if spec.note:
            B.set_notes(slide, spec.note)

    prs.save(path)
    return path


# =============================================================================
# 検証
# =============================================================================

def verify(B, path, specs, draft, with_logo=True):
    ok = True
    lines = []

    def check(label, cond, detail=""):
        nonlocal ok
        ok = ok and bool(cond)
        lines.append(f"  [{'OK' if cond else 'NG'}] {label}{(' — ' + detail) if detail else ''}")

    with zipfile.ZipFile(path) as z:
        names = z.namelist()
        check("zip 健全性", z.testzip() is None)
        theme = z.read("ppt/theme/theme1.xml").decode("utf-8")
        check(f"theme1.xml accent1={ACCENT1}",
              f'<a:accent1><a:srgbClr val="{ACCENT1}"/></a:accent1>' in theme)
        for key, val in B.CONFIG["colors"].items():
            if key != "accent1":
                check(f"theme1.xml {key}={val}", f'<a:{key}><a:srgbClr val="{val}"/></a:{key}>' in theme)
        svgs = {"ppt/media/kotonara-logo.svg", "ppt/media/kotonara-logo-white.svg"}
        if with_logo:
            check("SVG ロゴ 2点（緑・白）を同梱", svgs <= set(names))
        else:
            check("ロゴなし: 画像を同梱していない",
                  not [n for n in names if n.startswith("ppt/media/")])

        direct = []
        for n in names:
            if re.match(r"ppt/(slideMasters|slideLayouts|slides)/[^/]+\.xml$", n):
                if "srgbClr" in z.read(n).decode("utf-8"):
                    direct.append(n)
        check("マスター・レイアウト・スライドに色の直接指定なし", not direct, str(direct))

        slide_pics = [n for n in names if re.match(r"ppt/slides/slide\d+\.xml$", n)
                      and "<p:pic>" in z.read(n).decode("utf-8")]
        check("ロゴはスライドではなくレイアウトにある", not slide_pics, str(slide_pics))

    prs = Presentation(path)
    n = 1 if with_logo else 0
    want = {"表紙（タイトル）": n, "タイトルとコンテンツ": n, "表スライド": n,
            "セクション見出し": n, "クロージング（反転）": 0}
    for layout in prs.slide_layouts:
        pics = [sh for sh in layout.shapes if sh.shape_type == 13]
        check(f"レイアウト『{layout.name}』のロゴ {want.get(layout.name, 0)} 点",
              len(pics) == want.get(layout.name, 0), f"実際={len(pics)}")

    check(f"スライド数 = md の枚数 {len(specs)}", len(prs.slides) == len(specs))
    titles = [s.shapes.title.text_frame.text for s in prs.slides]
    md_titles = [B.BOLD_RE.sub(lambda m: m.group(0)[2:-2], s.title) for s in specs]
    check("タイトルが md と一致（`**` は外して比較）", titles == md_titles)
    text = B._all_text(prs)
    check("`**` 残骸なし", "**" not in text)
    check("`<!--` 残骸なし", "<!--" not in text and "-->" not in text)

    findings = B._placeholder_findings(prs)
    if not draft:
        check("{{...}} 残存なし", not findings,
              "; ".join(f"スライド{i}: {', '.join(t)}" for i, t in findings))

    print("\n".join(lines))
    if draft and findings:
        print("\n残っている {{...}}（--draft のためエラーにしていません）")
        for i, tokens in findings:
            for t in tokens:
                print(f"  スライド{i}: {t}")
    print("\n検証結果: " + ("ALL OK" if ok else "*** 失敗あり ***"))
    return ok


def build_one(md, out, footer="", with_logo=True, draft=False, doc_title=None):
    B = load_base(footer)
    specs, layouts = parse(B, md)
    print(f"md から {len(specs)} スライドを読み込みました: {md}")
    build_deck(B, specs, layouts, out, with_logo=with_logo, doc_title=doc_title)
    print(f"生成: {out}")
    print("-" * 60)
    return verify(B, out, specs, draft, with_logo=with_logo)


def build_templates(main_out=TEMPLATE_MAIN, nologo_out=TEMPLATE_NOLOGO_1):
    """メイン（ロゴ入り）とロゴなし版1を kotonara-brand-template.md から作る。"""
    ok = build_one(TEMPLATE_MD, main_out, with_logo=True,
                   doc_title="ことならテック ブランドテンプレート")
    print()
    ok = build_one(TEMPLATE_MD, nologo_out, with_logo=False,
                   doc_title="ことならテック ブランドテンプレート（ロゴなし版1）") and ok
    return ok


def main(argv=None):
    ap = argparse.ArgumentParser(description="ことならブランドテーマで md を PPTX にする")
    ap.add_argument("md", nargs="?")
    ap.add_argument("-o", "--out")
    ap.add_argument("--footer", default="", help="フッターの文言（既定は空欄）")
    ap.add_argument("--no-logo", action="store_true", help="ロゴを置かない（ロゴなし版1のテーマ）")
    ap.add_argument("--draft", action="store_true", help="{{...}} の残存を許す")
    ap.add_argument("--build-templates", action="store_true",
                    help="雛形2種（メインとロゴなし版1）をリポジトリ直下に作り直す")
    args = ap.parse_args(argv)

    if args.build_templates:
        return 0 if build_templates() else 1
    if not args.md:
        ap.error("md を指定してください（雛形を作り直すときは --build-templates）")
    out = args.out or os.path.splitext(args.md)[0] + "-kotonara.pptx"
    ok = build_one(args.md, out, footer=args.footer, with_logo=not args.no_logo, draft=args.draft)
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main())
