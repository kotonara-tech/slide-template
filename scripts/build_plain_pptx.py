# -*- coding: utf-8 -*-
"""
build_plain_pptx.py

確定済みMarkdownスライド原稿を、python-pptxで装飾最小限の
プレーンなPPTXに変換するスクリプト。

使い方:
    python build_plain_pptx.py [入力mdパス] [出力pptxパス]
引数を省略した場合は DEFAULT_MD_PATH / DEFAULT_OUT_PATH を使用する。

再実行可能: 何度実行しても同じ入力から同じ出力を生成する（上書き）。
"""

from __future__ import annotations

import re
import sys

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement

# ----------------------------------------------------------------------------
# 既定パス
# ----------------------------------------------------------------------------
DEFAULT_MD_PATH = (
    r"C:\Users\Ryuto\OneDrive\ドキュメント\karabiner\kototech"
    r"\kototech-aug-hackathon-kickoff-slides.md"
)
DEFAULT_OUT_PATH = (
    r"C:\Users\Ryuto\OneDrive\ドキュメント\karabiner\kototech"
    r"\kototech-aug-hackathon-kickoff-plain.pptx"
)

# ----------------------------------------------------------------------------
# デザイン定数
# ----------------------------------------------------------------------------
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

FONT_NAME = "游ゴシック"
TEXT_COLOR = RGBColor(0x33, 0x33, 0x33)
LINE_COLOR = "CCCCCC"          # タイトル下の罫線
TABLE_BORDER_COLOR = "BFBFBF"  # 表の罫線（控えめ）
TABLE_HEADER_FILL = RGBColor(0xE9, 0xE9, 0xE9)  # 薄グレー
TABLE_BODY_FILL = RGBColor(0xFF, 0xFF, 0xFF)

# 余白・レイアウト
MARGIN_L = Inches(0.7)
MARGIN_R = Inches(0.7)
CONTENT_W = SLIDE_W - MARGIN_L - MARGIN_R

TITLE_TOP = Inches(0.45)
TITLE_H = Inches(0.85)
RULE_TOP = Inches(1.32)
BODY_TOP = Inches(1.60)
BODY_BOTTOM_MARGIN = Inches(0.45)
BODY_H = SLIDE_H - BODY_TOP - BODY_BOTTOM_MARGIN

TITLE_SIZE = 30
COVER_TITLE_SIZE = 34
COVER_SUBTITLE_SIZE = 18
BODY_SIZE = 20
TABLE_HEADER_SIZE = 16
TABLE_BODY_SIZE = 15

# 表の「No Style, No Grid」相当のテーブルスタイルGUID（プレーン化のため）
PLAIN_TABLE_STYLE_GUID = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"

BOLD_RE = re.compile(r"\*\*(.+?)\*\*")
NOTE_RE = re.compile(r"<!--\s*note:\s*(.*?)\s*-->", re.DOTALL)


# ----------------------------------------------------------------------------
# Markdown パース
# ----------------------------------------------------------------------------
def split_slides(md_text: str) -> list[str]:
    """行全体が '---' のみの行でスライドを分割する。"""
    lines = md_text.split("\n")
    blocks: list[list[str]] = []
    current: list[str] = []
    for line in lines:
        if line.strip() == "---":
            blocks.append(current)
            current = []
        else:
            current.append(line)
    blocks.append(current)
    texts = ["\n".join(b) for b in blocks]
    # 空ブロック（先頭・末尾の余白等）を除去
    return [t for t in texts if t.strip()]


def parse_block(block_text: str):
    """1スライド分のテキストから (title, subtitle, content_lines, note) を取り出す。"""
    m = NOTE_RE.search(block_text)
    note = m.group(1).strip() if m else ""
    body_text = block_text[: m.start()] + block_text[m.end():] if m else block_text

    lines = body_text.split("\n")
    idx = 0

    def skip_blank(i):
        while i < len(lines) and lines[i].strip() == "":
            i += 1
        return i

    idx = skip_blank(idx)
    title = None
    if idx < len(lines) and lines[idx].lstrip().startswith("# "):
        title = lines[idx].lstrip()[2:].strip()
        idx += 1

    save_idx = idx
    idx2 = skip_blank(idx)
    subtitle = None
    if idx2 < len(lines) and lines[idx2].lstrip().startswith("## "):
        subtitle = lines[idx2].lstrip()[3:].strip()
        idx = idx2 + 1
    else:
        idx = save_idx

    content_lines = lines[idx:]
    return title, subtitle, content_lines, note


def is_table_sep_row(cells: list[str]) -> bool:
    return all(re.match(r"^:?-+:?$", c.strip()) for c in cells if c.strip() != "") and len(cells) > 0


def parse_content(content_lines: list[str]) -> list[dict]:
    """本文行を要素リスト（bullet / table / para）に変換する。"""
    elements: list[dict] = []
    i = 0
    n = len(content_lines)
    while i < n:
        line = content_lines[i]
        if line.strip() == "":
            i += 1
            continue
        stripped = line.lstrip()
        if stripped.startswith("|"):
            table_lines = []
            while i < n and content_lines[i].lstrip().startswith("|"):
                table_lines.append(content_lines[i].strip())
                i += 1
            rows = []
            for tl in table_lines:
                cells = [c.strip() for c in tl.strip("|").split("|")]
                if is_table_sep_row(cells):
                    continue
                rows.append(cells)
            if rows:
                elements.append({"type": "table", "rows": rows})
            continue
        m = re.match(r"^(\s*)-\s+(.*)$", line)
        if m:
            indent = len(m.group(1))
            level = 1 if indent >= 2 else 0
            text = m.group(2).strip()
            elements.append({"type": "bullet", "level": level, "text": text})
            i += 1
            continue
        # それ以外は段落として扱う
        elements.append({"type": "para", "text": line.strip()})
        i += 1
    return elements


def split_bold_runs(text: str):
    """'**bold**' を含むテキストを (chunk, is_bold) のリストへ分割する。"""
    parts = []
    last = 0
    for m in BOLD_RE.finditer(text):
        if m.start() > last:
            parts.append((text[last:m.start()], False))
        parts.append((m.group(1), True))
        last = m.end()
    if last < len(text):
        parts.append((text[last:], False))
    if not parts:
        parts = [(text, False)]
    return parts


def strip_bold_markers(text: str) -> str:
    return BOLD_RE.sub(lambda m: m.group(1), text)


# ----------------------------------------------------------------------------
# フォント適用ヘルパー（東アジア言語フォントをXMLレベルで確実に設定）
# ----------------------------------------------------------------------------
def set_east_asian_font(run, name: str = FONT_NAME) -> None:
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = OxmlElement("a:ea")
        rPr.append(ea)
    ea.set("typeface", name)
    # complex script (念のため) も同じフォントに揃える
    cs = rPr.find(qn("a:cs"))
    if cs is None:
        cs = OxmlElement("a:cs")
        rPr.append(cs)
    cs.set("typeface", name)


def apply_font(run, size=None, bold=None, color=TEXT_COLOR, name=FONT_NAME) -> None:
    run.font.name = name
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    run.font.color.rgb = color
    set_east_asian_font(run, name)


# ----------------------------------------------------------------------------
# 箇条書き（バレット記号）のXML設定
# ----------------------------------------------------------------------------
def set_paragraph_bullet(paragraph, char: str, level: int) -> None:
    pPr = paragraph._p.get_or_add_pPr()
    for tag in ("a:buChar", "a:buNone", "a:buAutoNum", "a:buFont"):
        el = pPr.find(qn(tag))
        if el is not None:
            pPr.remove(el)
    base_indent = 228600  # 0.25in
    step = 320000
    marL = base_indent + level * step
    indent = -228600
    pPr.set("marL", str(marL))
    pPr.set("indent", str(indent))
    buFont = OxmlElement("a:buFont")
    buFont.set("typeface", FONT_NAME)
    buChar = OxmlElement("a:buChar")
    buChar.set("char", char)
    pPr.append(buFont)
    pPr.append(buChar)


def clear_paragraph_bullet(paragraph) -> None:
    pPr = paragraph._p.get_or_add_pPr()
    for tag in ("a:buChar", "a:buAutoNum", "a:buFont"):
        el = pPr.find(qn(tag))
        if el is not None:
            pPr.remove(el)
    buNone = pPr.find(qn("a:buNone"))
    if buNone is None:
        buNone = OxmlElement("a:buNone")
        pPr.append(buNone)


# ----------------------------------------------------------------------------
# テキストボックス／段落ヘルパー
# ----------------------------------------------------------------------------
def add_textbox(slide, left, top, width, height):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    return box, tf


def add_run_text(paragraph, text, size, bold_base, color=TEXT_COLOR):
    for chunk, is_bold in split_bold_runs(text):
        if chunk == "":
            continue
        run = paragraph.add_run()
        run.text = chunk
        apply_font(run, size=size, bold=(bold_base or is_bold), color=color)


def add_title_box(slide, text, *, left, top, width, height, size, align, bold=True):
    box, tf = add_textbox(slide, left, top, width, height)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = align
    add_run_text(p, text, size, bold_base=bold)
    return box


def add_rule_line(slide, left, top, width):
    line = slide.shapes.add_connector(1, left, top, left + width, top)  # 1 = straight
    line.line.color.rgb = RGBColor.from_string(LINE_COLOR)
    line.line.width = Pt(0.75)
    return line


def add_body_content(slide, elements, *, left, top, width, height, base_size=BODY_SIZE):
    box, tf = add_textbox(slide, left, top, width, height)
    first = True
    for el in elements:
        if el["type"] == "bullet":
            if first and not tf.paragraphs[0].runs:
                p = tf.paragraphs[0]
            else:
                p = tf.add_paragraph()
            first = False
            p.alignment = PP_ALIGN.LEFT
            p.line_spacing = 1.15
            p.space_after = Pt(10)
            char = "・" if el["level"] == 0 else "-"
            add_run_text(p, el["text"], base_size, bold_base=False)
            set_paragraph_bullet(p, char, el["level"])
        else:  # para
            if first and not tf.paragraphs[0].runs:
                p = tf.paragraphs[0]
            else:
                p = tf.add_paragraph()
            first = False
            p.alignment = PP_ALIGN.LEFT
            p.line_spacing = 1.15
            p.space_after = Pt(10)
            add_run_text(p, el["text"], base_size, bold_base=False)
            clear_paragraph_bullet(p)
    return box


# ----------------------------------------------------------------------------
# 表ヘルパー
# ----------------------------------------------------------------------------
def compute_col_widths(rows, total_width_emu, min_ratio=0.14):
    n_cols = len(rows[0])
    max_lens = [1] * n_cols
    for row in rows:
        for i, cell in enumerate(row):
            plain = strip_bold_markers(cell)
            max_lens[i] = max(max_lens[i], len(plain))
    total_len = sum(max_lens)
    ratios = [max(l / total_len, min_ratio) for l in max_lens]
    s = sum(ratios)
    ratios = [r / s for r in ratios]
    widths = [int(total_width_emu * r) for r in ratios]
    # 端数をラスト列で吸収
    diff = total_width_emu - sum(widths)
    widths[-1] += diff
    return widths


def set_cell_borders(cell, color=TABLE_BORDER_COLOR, width_pt=0.75):
    tcPr = cell._tc.get_or_add_tcPr()
    w = int(Pt(width_pt))
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        existing = tcPr.find(qn(tag))
        if existing is not None:
            tcPr.remove(existing)
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        ln = OxmlElement(tag)
        ln.set("w", str(w))
        ln.set("cap", "flat")
        ln.set("cmpd", "sng")
        ln.set("algn", "ctr")
        fill = OxmlElement("a:solidFill")
        clr = OxmlElement("a:srgbClr")
        clr.set("val", color)
        fill.append(clr)
        ln.append(fill)
        tcPr.append(ln)


def add_table_slide(slide, rows, *, left, top, width, height):
    n_rows = len(rows)
    n_cols = len(rows[0])
    graphic_frame = slide.shapes.add_table(n_rows, n_cols, left, top, width, height)
    table = graphic_frame.table
    table.first_row = False
    table.horz_banding = False

    # テーブルスタイルをプレーンなものに変更
    tblPr = table._tbl.find(qn("a:tblPr"))
    if tblPr is not None:
        style_el = tblPr.find(qn("a:tableStyleId"))
        if style_el is not None:
            style_el.text = PLAIN_TABLE_STYLE_GUID

    col_widths = compute_col_widths(rows, width)
    for c, w in enumerate(col_widths):
        table.columns[c].width = Emu(w)

    row_h = int(height / n_rows)
    for r in range(n_rows):
        table.rows[r].height = Emu(row_h)

    for r, row_data in enumerate(rows):
        is_header = r == 0
        for c, cell_text in enumerate(row_data):
            cell = table.cell(r, c)
            cell.margin_left = Emu(91440)
            cell.margin_right = Emu(91440)
            cell.margin_top = Emu(36576)
            cell.margin_bottom = Emu(36576)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE

            set_cell_borders(cell)

            cell.fill.solid()
            cell.fill.fore_color.rgb = TABLE_HEADER_FILL if is_header else TABLE_BODY_FILL

            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT
            add_run_text(
                p,
                cell_text,
                TABLE_HEADER_SIZE if is_header else TABLE_BODY_SIZE,
                bold_base=is_header,
            )
    return graphic_frame


# ----------------------------------------------------------------------------
# スライド構築
# ----------------------------------------------------------------------------
def build_cover_slide(prs, title, subtitle, elements, note):
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    add_title_box(
        slide,
        title,
        left=Inches(0.9),
        top=Inches(1.5),
        width=Inches(11.533),
        height=Inches(1.6),
        size=COVER_TITLE_SIZE,
        align=PP_ALIGN.CENTER,
        bold=True,
    )

    if subtitle:
        add_title_box(
            slide,
            subtitle,
            left=Inches(0.9),
            top=Inches(3.15),
            width=Inches(11.533),
            height=Inches(0.9),
            size=COVER_SUBTITLE_SIZE,
            align=PP_ALIGN.CENTER,
            bold=False,
        )

    if elements:
        add_body_content(
            slide,
            elements,
            left=Inches(1.6),
            top=Inches(4.35),
            width=Inches(10.133),
            height=Inches(2.6),
            base_size=BODY_SIZE,
        )

    set_notes(slide, note)
    return slide


def build_normal_slide(prs, title, elements, note):
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    add_title_box(
        slide,
        title,
        left=MARGIN_L,
        top=TITLE_TOP,
        width=CONTENT_W,
        height=TITLE_H,
        size=TITLE_SIZE,
        align=PP_ALIGN.LEFT,
        bold=True,
    )
    add_rule_line(slide, MARGIN_L, RULE_TOP, CONTENT_W)

    tables = [e for e in elements if e["type"] == "table"]
    non_table = [e for e in elements if e["type"] != "table"]

    cursor_top = BODY_TOP
    remaining_h = BODY_H

    if non_table:
        # 非表要素がある場合、必要な高さを簡易見積り
        approx_h = Inches(0.55) * sum(1 for e in non_table)
        text_h = min(remaining_h, max(approx_h, Inches(1.0))) if tables else remaining_h
        add_body_content(
            slide,
            non_table,
            left=Inches(0.9),
            top=cursor_top,
            width=Inches(11.533),
            height=text_h,
            base_size=BODY_SIZE,
        )
        if tables:
            cursor_top = cursor_top + text_h + Inches(0.15)
            remaining_h = SLIDE_H - cursor_top - BODY_BOTTOM_MARGIN

    for tbl_el in tables:
        add_table_slide(
            slide,
            tbl_el["rows"],
            left=Inches(0.9),
            top=cursor_top,
            width=Inches(11.533),
            height=remaining_h,
        )

    set_notes(slide, note)
    return slide


def set_notes(slide, note_text: str) -> None:
    notes_slide = slide.notes_slide
    tf = notes_slide.notes_text_frame
    tf.text = note_text


# ----------------------------------------------------------------------------
# メイン
# ----------------------------------------------------------------------------
def main():
    md_path = sys.argv[1] if len(sys.argv) > 1 else DEFAULT_MD_PATH
    out_path = sys.argv[2] if len(sys.argv) > 2 else DEFAULT_OUT_PATH

    with open(md_path, "r", encoding="utf-8") as f:
        md_text = f.read()

    blocks = split_slides(md_text)

    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    for i, block in enumerate(blocks):
        title, subtitle, content_lines, note = parse_block(block)
        elements = parse_content(content_lines)
        if title is None:
            raise ValueError(f"スライド{i + 1}にタイトル（'# '見出し）が見つかりません")
        if i == 0:
            build_cover_slide(prs, title, subtitle, elements, note)
        else:
            build_normal_slide(prs, title, elements, note)

    prs.save(out_path)
    print(f"生成完了: {out_path} （スライド数: {len(blocks)}）")


if __name__ == "__main__":
    main()
