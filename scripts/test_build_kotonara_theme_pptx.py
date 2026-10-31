# -*- coding: utf-8 -*-
"""build_kotonara_theme_pptx.py の受け入れテスト。

実行（リポジトリルートで）:
  python -m pytest -q scripts/test_build_kotonara_theme_pptx.py
"""
import os
import sys
import zipfile

from pptx import Presentation

SCRIPTS = os.path.dirname(os.path.abspath(__file__))
REPO_ROOT = os.path.dirname(SCRIPTS)
sys.path.insert(0, SCRIPTS)

import build_kotonara_theme_pptx as K  # noqa: E402

LAYOUTS = ["表紙（タイトル）", "タイトルとコンテンツ", "表スライド", "セクション見出し", "クロージング（反転）"]


def build(tmp_path, md, *extra):
    out = str(tmp_path / "out.pptx")
    code = K.main([os.path.join(REPO_ROOT, md), "-o", out, *extra])
    return code, out


def logo_counts(prs):
    return {l.name: sum(1 for sh in l.shapes if sh.shape_type == 13) for l in prs.slide_layouts}


def cover_catch_size(path):
    import re
    with zipfile.ZipFile(path) as z:
        x = z.read("ppt/slideLayouts/slideLayout1.xml").decode("utf-8")
    i = x.find("Cover Band")
    return re.search(r'sz="(\d+)"', x[i:x.find("</p:sp>", i)]).group(1)


def theme_xml(path):
    with zipfile.ZipFile(path) as z:
        return z.read("ppt/theme/theme1.xml").decode("utf-8")


def test_build_templates_main_and_nologo_1(tmp_path):
    main, nologo = str(tmp_path / "main.pptx"), str(tmp_path / "nologo1.pptx")
    assert K.build_templates(main, nologo)
    for path in (main, nologo):
        prs = Presentation(path)
        assert [s.slide_layout.name for s in prs.slides] == LAYOUTS
        assert '<a:accent1><a:srgbClr val="4D9622"/></a:accent1>' in theme_xml(path)
    assert logo_counts(Presentation(main)) == {
        "表紙（タイトル）": 1, "タイトルとコンテンツ": 1, "表スライド": 1,
        "セクション見出し": 1, "クロージング（反転）": 0}
    assert set(logo_counts(Presentation(nologo)).values()) == {0}


def test_logos_live_in_layouts_not_slides(tmp_path):
    _, out = build(tmp_path, "kotonara-brand-template.md")
    for slide in Presentation(out).slides:
        assert not any(sh.shape_type == 13 for sh in slide.shapes)


def test_no_logo_option(tmp_path):
    code, out = build(tmp_path, "kotonara-brand-template.md", "--no-logo")
    assert code == 0
    with zipfile.ZipFile(out) as z:
        assert not [n for n in z.namelist() if n.startswith("ppt/media/")]


def test_footer_empty_by_default_and_option(tmp_path):
    _, out = build(tmp_path, "kotonara-brand-template.md")
    texts = [sh.text_frame.text for sh in Presentation(out).slides[1].placeholders if sh.has_text_frame]
    assert "フッターの確認" not in texts
    _, out = build(tmp_path, "kotonara-brand-template.md", "--footer", "フッターの確認")
    texts = [sh.text_frame.text for sh in Presentation(out).slides[1].placeholders if sh.has_text_frame]
    assert "フッターの確認" in texts


def test_committed_templates_are_as_expected():
    main = os.path.join(REPO_ROOT, "kotonara-brand-template.pptx")
    nologo1 = os.path.join(REPO_ROOT, "kotonara-brand-template-nologo-1.pptx")
    nologo2 = os.path.join(REPO_ROOT, "kotonara-brand-template-nologo-2.pptx")
    assert '<a:srgbClr val="4D9622"/>' in theme_xml(main)
    assert '<a:srgbClr val="4D9622"/>' in theme_xml(nologo1)
    assert '<a:accent1><a:srgbClr val="1A832A"/></a:accent1>' in theme_xml(nologo2)
    assert logo_counts(Presentation(main))["タイトルとコンテンツ"] == 1
    assert set(logo_counts(Presentation(nologo1)).values()) == {0}
    assert set(logo_counts(Presentation(nologo2)).values()) == {0}
    for path in (main, nologo1, nologo2):
        assert cover_catch_size(path) == "1300"   # キャッチコピーは3種とも 13pt


def test_calendar_sample(tmp_path):
    code, out = build(tmp_path, "examples/kotonara-calendar-sample.md")
    assert code == 0
    table = [sh for sh in Presentation(out).slides[1].shapes if sh.has_table][0].table
    assert len(table.columns) == 8
    assert table.cell(3, 6).text_frame.text == "19\n開催当日"


def test_placeholder_fails_without_draft(tmp_path):
    md = tmp_path / "p.md"
    md.write_text("# 表紙\n\n## {{サブタイトル}}\n\n---\n\n# おわり\n", encoding="utf-8")
    out = str(tmp_path / "p.pptx")
    assert K.main([str(md), "-o", out]) == 1
    assert K.main([str(md), "-o", out, "--draft"]) == 0


def write_md(tmp_path, text):
    md = tmp_path / "t.md"
    md.write_text(text, encoding="utf-8")
    return str(md), str(tmp_path / "t.pptx")


CAL_HEAD = "| 日 | 月 | 火 | 水 | 木 | 金 | 土 | この週のこと |\n| --- | --- | --- | --- | --- | --- | --- | --- |\n"


def cal_md(body):
    return f"# 表紙\n\n---\n\n# 予定\n\n{body}\n---\n\n# おわり\n"


def test_calendar_bad_shapes_fail_with_message(tmp_path):
    cases = [
        "| 日 | 月 | 火 | 水 | 木 | 金 | 土 |\n| --- | --- | --- | --- | --- | --- | --- |\n| 1 | 2 | 3 | 4 | 5 | 6 | 7 |\n",
        CAL_HEAD + "| 1 | 2 | 3 |\n",
        CAL_HEAD + "| 1 | 2 | 3 | 4 | 5 | 6 | 7 | a | b |\n",
        CAL_HEAD,
    ]
    for body in cases:
        md, out = write_md(tmp_path, cal_md(body))
        try:
            K.main([md, "-o", out])
        except SystemExit as e:
            assert "カレンダー表" in str(e.code)
        else:
            raise AssertionError(f"エラーにならない: {body!r}")


def test_calendar_fullwidth_space_label(tmp_path):
    md, out = write_md(tmp_path, cal_md(CAL_HEAD + "| 13 | 14 | 15 | 16 | 17 | 18 | 19　開催当日 | 当日 |\n"))
    assert K.main([md, "-o", out]) == 0
    table = [sh for sh in Presentation(out).slides[1].shapes if sh.has_table][0].table
    assert table.cell(1, 6).text_frame.text == "19\n開催当日"


def test_table_outside_table_layout_fails(tmp_path):
    for text in ("# 表紙\n\n---\n\n# おわり\n\n| a | b |\n| --- | --- |\n| 1 | 2 |\n",
                 "# 表紙\n\n---\n\n# 章\n\n<!-- layout: section -->\n\n| a | b |\n| --- | --- |\n| 1 | 2 |\n"
                 "\n---\n\n# おわり\n"):
        md, out = write_md(tmp_path, text)
        try:
            K.main([md, "-o", out])
        except SystemExit as e:
            assert "表は表スライドにしか置けません" in str(e.code)
        else:
            raise AssertionError("表が黙って消えた")


def test_bullets_on_table_slide_fail(tmp_path):
    md, out = write_md(tmp_path, "# 表紙\n\n---\n\n# 表\n\n- 消える箇条書き\n\n| a | b |\n| --- | --- |\n"
                                 "| 1 | 2 |\n\n---\n\n# おわり\n")
    try:
        K.main([md, "-o", out])
    except SystemExit as e:
        assert "箇条書き・本文を書けません" in str(e.code)
    else:
        raise AssertionError("箇条書きが黙って消えた")


def test_layout_only_block_fails(tmp_path):
    md, out = write_md(tmp_path, "# 表紙\n\n---\n\n<!-- layout: section -->\n\n---\n\n# 本文\n\n- a\n\n---\n\n# おわり\n")
    try:
        K.main([md, "-o", out])
    except SystemExit as e:
        assert "レイアウト指定だけ" in str(e.code)
    else:
        raise AssertionError("レイアウト指定だけの枚がエラーにならない")


def test_bold_title_passes_verify(tmp_path):
    md, out = write_md(tmp_path, "# 表紙\n\n---\n\n# **強調**の見出し\n\n- a\n\n---\n\n# おわり\n")
    assert K.main([md, "-o", out]) == 0


def test_title_does_not_run_under_logo(tmp_path):
    _, out = build(tmp_path, "kotonara-brand-template.md")
    prs = Presentation(out)
    for name in ("タイトルとコンテンツ", "表スライド"):
        layout = [l for l in prs.slide_layouts if l.name == name][0]
        logo = [sh for sh in layout.shapes if sh.shape_type == 13][0]
        title = [sh for sh in layout.placeholders if sh.placeholder_format.idx == 0][0]
        assert title.left + title.width <= logo.left
