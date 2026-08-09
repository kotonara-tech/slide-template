# -*- coding: utf-8 -*-
"""
verify_pptx.py

build_plain_pptx.py が生成したPPTXを検証する。
- スライド数 = 19
- 全19スライドにスピーカーノートが存在し空でない
- 3枚（アジェンダ/Day1/Day2）に表が存在する
- 本文に '<!-- note:' や '**' の残骸が混入していない
- タイトルテキストがmdの '# ' 19本と一致
"""

from __future__ import annotations

import re
import sys

from pptx import Presentation
from pptx.util import Emu

DEFAULT_MD_PATH = (
    r"C:\Users\Ryuto\OneDrive\ドキュメント\karabiner\kototech"
    r"\kototech-aug-hackathon-kickoff-slides.md"
)
DEFAULT_PPTX_PATH = (
    r"C:\Users\Ryuto\OneDrive\ドキュメント\karabiner\kototech"
    r"\kototech-aug-hackathon-kickoff-plain.pptx"
)


def extract_md_titles(md_text: str) -> list[str]:
    titles = []
    for line in md_text.split("\n"):
        if line.startswith("# ") and not line.startswith("## "):
            titles.append(line[2:].strip())
    return titles


def iter_shape_texts(shapes):
    for shape in shapes:
        if shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                text = "".join(r.text for r in p.runs)
                if text:
                    yield text
        if shape.has_table:
            for row in shape.table.rows:
                for cell in row.cells:
                    for p in cell.text_frame.paragraphs:
                        text = "".join(r.text for r in p.runs)
                        if text:
                            yield text


def main():
    md_path = sys.argv[1] if len(sys.argv) > 1 else DEFAULT_MD_PATH
    pptx_path = sys.argv[2] if len(sys.argv) > 2 else DEFAULT_PPTX_PATH

    with open(md_path, "r", encoding="utf-8") as f:
        md_text = f.read()
    md_titles = extract_md_titles(md_text)

    prs = Presentation(pptx_path)
    slides = list(prs.slides)

    errors = []
    warnings = []

    # 1) スライド数 = 19
    if len(slides) != 19:
        errors.append(f"スライド数が19ではありません: {len(slides)}")
    else:
        print(f"OK: スライド数 = {len(slides)}")

    # 2) 全スライドにスピーカーノートが存在し空でない
    empty_notes = []
    for i, slide in enumerate(slides, start=1):
        if not slide.has_notes_slide:
            empty_notes.append(i)
            continue
        note_text = slide.notes_slide.notes_text_frame.text.strip()
        if not note_text:
            empty_notes.append(i)
    if empty_notes:
        errors.append(f"スピーカーノートが空/存在しないスライド: {empty_notes}")
    else:
        print("OK: 全19スライドにスピーカーノートあり（空でない）")

    # 3) 3枚に表が存在する（アジェンダ/Day1/Day2）
    table_slide_indices = []
    for i, slide in enumerate(slides, start=1):
        has_table = any(shape.has_table for shape in slide.shapes)
        if has_table:
            table_slide_indices.append(i)
    if len(table_slide_indices) != 3:
        errors.append(f"表を含むスライド数が3ではありません: {table_slide_indices}")
    else:
        print(f"OK: 表を含むスライド = {table_slide_indices} （3枚）")

    # タイトルで具体的にアジェンダ/Day1/Day2であることも確認
    expected_table_title_keywords = ["アジェンダ", "Day1", "Day2"]
    title_by_slide = {}
    for i, slide in enumerate(slides, start=1):
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                title_by_slide[i] = shape.text_frame.text.strip()
                break
    for idx, kw in zip(table_slide_indices, expected_table_title_keywords):
        t = title_by_slide.get(idx, "")
        if kw not in t:
            warnings.append(f"表スライド{idx}のタイトルに'{kw}'が含まれません: {t!r}")

    # 4) 本文に '<!-- note:' や '**' の残骸が混入していない
    contaminated = []
    for i, slide in enumerate(slides, start=1):
        for text in iter_shape_texts(slide.shapes):
            if "<!-- note:" in text or "-->" in text:
                contaminated.append((i, "note-comment", text))
            if "**" in text:
                contaminated.append((i, "bold-marker", text))
    if contaminated:
        errors.append(f"本文にnoteコメント/**の残骸あり: {contaminated}")
    else:
        print("OK: 本文に '<!-- note:' や '**' の残骸なし")

    # スピーカーノート側にnoteマーカーが残っていないかも確認（意図的に'note:'部分のみ抽出しているはず）
    note_contaminated = []
    for i, slide in enumerate(slides, start=1):
        if slide.has_notes_slide:
            note_text = slide.notes_slide.notes_text_frame.text
            if "<!--" in note_text or "-->" in note_text:
                note_contaminated.append(i)
    if note_contaminated:
        warnings.append(f"スピーカーノートにHTMLコメント記号が残存: {note_contaminated}")

    # 5) タイトルテキストがmdの '# ' 19本と一致
    pptx_titles = []
    for i, slide in enumerate(slides, start=1):
        # タイトルは各スライドの最初のテキストフレーム（表以外）
        title_text = None
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                title_text = shape.text_frame.text.strip()
                break
        pptx_titles.append(title_text)

    if len(md_titles) != 19:
        errors.append(f"md側の'# 'タイトル数が19ではありません: {len(md_titles)}")

    mismatches = []
    for i, (md_t, pptx_t) in enumerate(zip(md_titles, pptx_titles), start=1):
        if md_t != pptx_t:
            mismatches.append((i, md_t, pptx_t))
    if mismatches:
        errors.append(f"タイトル不一致: {mismatches}")
    else:
        print("OK: 全19スライドのタイトルがmdの'# '見出しと一致")

    # 表がスライド幅からはみ出していないかの簡易チェック
    overflow = []
    for i, slide in enumerate(slides, start=1):
        for shape in slide.shapes:
            if shape.has_table:
                right = shape.left + shape.width
                bottom = shape.top + shape.height
                if right > prs.slide_width or bottom > prs.slide_height:
                    overflow.append(i)
    if overflow:
        warnings.append(f"表がスライド範囲をはみ出している可能性: {overflow}")
    else:
        print("OK: 表はスライド範囲内に収まっている")

    print()
    if warnings:
        print("=== 警告 ===")
        for w in warnings:
            print(f"- {w}")

    if errors:
        print("=== エラー ===")
        for e in errors:
            print(f"- {e}")
        print("\n検証: 失敗")
        sys.exit(1)
    else:
        print("検証: 全項目パス")


if __name__ == "__main__":
    main()
